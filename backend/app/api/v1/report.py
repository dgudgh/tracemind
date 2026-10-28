import io
import html
import json
import logging
from datetime import datetime
from types import SimpleNamespace

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.itsm.models import ItsmTicket
from app.models.ai_report import AIReport
from app.models.alert_event import DBAlertEvent
from app.models.backup_operation import BackupOperation
from app.models.backup_platform import BackupExternalRun
from app.models.data_governance import DBArchiveJob
from app.models.dbcheck_report import DBCheckReport
from app.models.db_instance import DBInstance
from app.models.oracle_awr import AwrReport
from app.models.perf_metric import DBPerfMetric
from app.models.slow_sql import DBSlowSQL
from app.schemas.report import ReportGenerateRequest, ReportResponse
from app.services import instances as instance_service
from app.services.dbcheck import run_dbcheck
from app.services.inspection import dispatch as inspection_dispatch
from app.services.inspection import report_artifacts, report_builder, result_writer, run_readback
from app.services.reports import generate_markdown_report

logger = logging.getLogger("app.api.report")

router = APIRouter(prefix="/reports", tags=["巡检报告中心"])


@router.get("")
def list_reports(db: Session = Depends(get_db)):
    rows = db.query(DBCheckReport).order_by(DBCheckReport.id.desc()).limit(50).all()
    return {"items": [
        {
            "id": r.id,
            "instance_id": r.instance_id,
            "instance_name": r.instance_name,
            "total": r.total,
            "success": r.success,
            "failed": r.failed,
            "create_time": r.create_time.isoformat(sep=" ") if r.create_time else None,
        }
        for r in rows
    ]}


@router.post("/generate", response_model=ReportResponse)
def generate_report(data: ReportGenerateRequest, db: Session = Depends(get_db)):
    content = generate_markdown_report(db, data.task_id, data.instance_id)
    return ReportResponse(
        title="TraceMind AI 智能运维平台 巡检报告",
        format=data.output_format,
        content=content,
        generated_at=datetime.now(),
    )


class DBCheckRequest(BaseModel):
    instance_id: int


def _dt(value):
    return value.isoformat(sep=" ") if value else None


def _instance_lookup(db: Session) -> dict[int, DBInstance]:
    return {row.id: row for row in db.query(DBInstance).all()}


def _inst_name(instances: dict[int, DBInstance], instance_id: int | None) -> str:
    if not instance_id:
        return "-"
    inst = instances.get(instance_id)
    return inst.instance_name if inst else str(instance_id)


def _limit_query(query, limit: int):
    return query.limit(max(1, min(limit, 500))).all()


def _collect_unified_reports(db: Session, report_type: str = "all", limit: int = 100) -> list[dict]:
    instances = _instance_lookup(db)
    items: list[dict] = []

    def enabled(kind: str) -> bool:
        return report_type in ("all", kind)

    if enabled("inspection"):
        # The catalogue reads the run table, not the result rows. Aggregating from rows meant a
        # failed or cancelled run - which has no rows - was absent from the catalogue entirely,
        # so "the report failed" and "the run never happened" looked the same.
        for item in run_readback.run_sources(db, limit):
            items.append(report_artifacts.source_item(item))

    if enabled("awr"):
        rows = _limit_query(db.query(AwrReport).order_by(AwrReport.id.desc()), limit)
        for row in rows:
            items.append({
                "id": f"awr:{row.id}",
                "report_type": "awr",
                "source_id": row.id,
                "instance_id": row.instance_id,
                "instance_name": _inst_name(instances, row.instance_id),
                "db_type": "oracle",
                "title": row.file_name or f"AWR {row.begin_snap}-{row.end_snap}",
                "status": row.status,
                "risk_level": "warning" if row.performance_score is not None and row.performance_score < 75 else "normal",
                "score": row.performance_score,
                "summary": f"Snap {row.begin_snap}-{row.end_snap}，DB {row.db_name or '-'}",
                "create_time": _dt(row.create_time),
                "export_url": f"/api/v1/oracle/{row.instance_id}/awr/reports/{row.id}/export?format=md",
            })

    if enabled("performance"):
        rows = (
            db.query(DBPerfMetric)
            .order_by(DBPerfMetric.capture_time.desc(), DBPerfMetric.id.desc())
            .limit(max(20, min(limit * 20, 5000)))
            .all()
        )
        by_instance: dict[int, list[DBPerfMetric]] = {}
        for row in rows:
            by_instance.setdefault(row.instance_id, []).append(row)
        for instance_id, samples in list(by_instance.items())[:limit]:
            latest_by_metric: dict[str, DBPerfMetric] = {}
            for sample in samples:
                latest_by_metric.setdefault(sample.metric_name, sample)
            latest = max(samples, key=lambda sample: (sample.capture_time, sample.id))
            metric_values = {
                name: float(sample.value) if sample.value is not None else None
                for name, sample in sorted(latest_by_metric.items())
            }
            cpu_value = metric_values.get("cpu_util")
            risk_level = "warning" if cpu_value is not None and cpu_value >= 85 else "normal"
            summary = "，".join(
                f"{name}={value:g}" if value is not None else f"{name}=-"
                for name, value in metric_values.items()
            ) or "暂无有效性能指标"
            items.append({
                "id": f"performance:{latest.id}",
                "report_type": "performance",
                "source_id": f"instance-{instance_id}-sample-{latest.id}",
                "instance_id": instance_id,
                "instance_name": _inst_name(instances, instance_id),
                "db_type": latest.db_type,
                "title": f"性能基线快照 · {_inst_name(instances, instance_id)}",
                "status": "captured",
                "risk_level": risk_level,
                "score": None,
                "summary": summary,
                "create_time": _dt(latest.capture_time),
                "export_url": None,
            })

    if enabled("slow_sql"):
        rows = _limit_query(db.query(DBSlowSQL).order_by(DBSlowSQL.id.desc()), limit)
        for row in rows:
            items.append({
                "id": f"slow_sql:{row.id}",
                "report_type": "slow_sql",
                "source_id": row.id,
                "instance_id": row.instance_id,
                "instance_name": _inst_name(instances, row.instance_id),
                "db_type": row.db_type,
                "title": f"慢 SQL {row.sql_id or row.id}",
                "status": "captured",
                "risk_level": row.risk_level,
                "score": None,
                "summary": "已采集 SQL 指标快照；本报告未执行 AI 或执行计划分析。" + (row.sql_text or "")[:160],
                "create_time": _dt(row.capture_time or row.create_time),
                "export_url": None,
                "slow_sql_detail": {
                    "sql_id": row.sql_id,
                    "schema_name": row.schema_name,
                    "sql_text": row.sql_text,
                    "capture_time": _dt(row.capture_time),
                    "create_time": _dt(row.create_time),
                    "metrics": {
                        "executions": row.executions,
                        "avg_elapsed_ms": row.avg_elapsed_ms,
                        "max_elapsed_ms": row.max_elapsed_ms,
                        "total_elapsed_ms": row.total_elapsed_ms,
                        "logical_reads": row.logical_reads,
                        "physical_reads": row.physical_reads,
                        "rows_processed": row.rows_processed,
                    },
                    "involved_tables": row.involved_tables,
                    "mes_module": row.mes_module,
                    "analysis_status": "not_run",
                    "plan_status": "not_run",
                    "statistics_scope": "legacy_cumulative_snapshot",
                    "sql_text_scope": "stored_text_unverified_completeness",
                    "legacy_notes": {
                        "verification_status": "unverified",
                        "ai_summary": row.ai_summary,
                        "optimize_suggestion": row.optimize_suggestion,
                    },
                },
            })

    if enabled("alert"):
        rows = _limit_query(db.query(DBAlertEvent).order_by(DBAlertEvent.last_seen.desc(), DBAlertEvent.id.desc()), limit)
        for row in rows:
            items.append({
                "id": f"alert:{row.id}",
                "report_type": "alert",
                "source_id": row.id,
                "instance_id": row.instance_id,
                "instance_name": _inst_name(instances, row.instance_id),
                "db_type": row.db_type,
                "title": row.title,
                "status": row.status,
                "risk_level": row.severity,
                "score": None,
                "summary": row.message,
                "create_time": _dt(row.last_seen or row.create_time),
                "export_url": None,
            })

    if enabled("ticket"):
        rows = _limit_query(db.query(ItsmTicket).order_by(ItsmTicket.id.desc()), limit)
        for row in rows:
            items.append({
                "id": f"ticket:{row.id}",
                "report_type": "ticket",
                "source_id": row.id,
                "instance_id": row.instance_id,
                "instance_name": row.instance_name or _inst_name(instances, row.instance_id),
                "db_type": row.db_type,
                "title": row.title,
                "status": row.mesguard_status,
                "risk_level": row.severity or row.priority,
                "score": None,
                "summary": row.description,
                "create_time": _dt(row.create_time),
                "export_url": None,
            })

    if enabled("governance"):
        rows = _limit_query(db.query(DBArchiveJob).order_by(DBArchiveJob.id.desc()), limit)
        for row in rows:
            items.append({
                "id": f"governance:{row.id}",
                "report_type": "governance",
                "source_id": row.id,
                "instance_id": row.instance_id,
                "instance_name": _inst_name(instances, row.instance_id),
                "db_type": row.db_type,
                "title": row.job_name,
                "status": row.status,
                "risk_level": "info",
                "score": None,
                "summary": f"{row.job_type}: {row.schema_name}.{row.table_name}",
                "create_time": _dt(row.create_time),
                "export_url": None,
            })

    if enabled("ai"):
        rows = _limit_query(db.query(AIReport).order_by(AIReport.id.desc()), limit)
        for row in rows:
            items.append({
                "id": f"ai:{row.id}",
                "report_type": "ai",
                "source_id": row.id,
                "instance_id": row.instance_id,
                "instance_name": _inst_name(instances, row.instance_id),
                "db_type": row.db_type,
                "title": row.title,
                "status": (row.diagnosis or {}).get("status", "generated") if isinstance(row.diagnosis, dict) else "generated",
                "risk_level": row.risk_level,
                "score": None,
                "summary": row.summary,
                "create_time": _dt(row.create_time),
                "export_url": None,
                **({"analysis_detail": row.diagnosis} if isinstance(row.diagnosis, dict) and row.diagnosis.get("kind") == "slow_sql_ai" else {}),
            })

    if enabled("backup"):
        rows = _limit_query(db.query(BackupOperation).order_by(BackupOperation.id.desc()), limit)
        for row in rows:
            failed = str(row.status or "").lower() in {"failed", "error", "cancelled"}
            items.append({
                "id": f"backup:operation:{row.id}",
                "report_type": "backup",
                "source_id": f"operation-{row.id}",
                "instance_id": row.instance_id,
                "instance_name": _inst_name(instances, row.instance_id),
                "db_type": row.engine,
                "title": f"{row.category} · {row.operation_no}",
                "status": row.status,
                "risk_level": "critical" if failed else row.risk_level or "normal",
                "score": None,
                "summary": (
                    f"{row.operation_type}，校验 {row.validation_status or '未执行'}，"
                    f"耗时 {row.duration_seconds if row.duration_seconds is not None else '-'} 秒"
                ),
                "create_time": _dt(row.finished_at or row.started_at or row.create_time),
                "export_url": None,
            })
        external_rows = _limit_query(
            db.query(BackupExternalRun).order_by(BackupExternalRun.finished_at.desc(), BackupExternalRun.id.desc()),
            limit,
        )
        for row in external_rows:
            failed = str(row.status or "").lower() in {"failed", "error", "cancelled"}
            items.append({
                "id": f"backup:external:{row.id}",
                "report_type": "backup",
                "source_id": f"external-{row.id}",
                "instance_id": row.instance_id,
                "instance_name": _inst_name(instances, row.instance_id),
                "db_type": getattr(instances.get(row.instance_id), "db_type", None),
                "title": f"外部备份运行 · {row.external_run_id}",
                "status": row.status,
                "risk_level": "critical" if failed else "normal",
                "score": None,
                "summary": row.error_message or f"外部任务 {row.external_job_id}，容量 {row.capacity_bytes or 0} 字节",
                "create_time": _dt(row.finished_at or row.started_at or row.synced_at),
                "export_url": None,
            })

    items.sort(key=lambda item: item.get("create_time") or "", reverse=True)
    return items[:limit]


@router.get("/unified")
def list_unified_reports(
    report_type: str = Query("all", pattern="^(all|inspection|awr|performance|slow_sql|alert|ticket|backup|governance|ai|rescue)$"),
    limit: int = Query(100, ge=1, le=500),
    db: Session = Depends(get_db),
):
    return {"items": _collect_unified_reports(db, report_type, limit)}


@router.get("/unified/stats")
def unified_report_stats(db: Session = Depends(get_db)):
    items = _collect_unified_reports(db, "all", 500)
    counts: dict[str, int] = {}
    risks: dict[str, int] = {}
    for item in items:
        counts[item["report_type"]] = counts.get(item["report_type"], 0) + 1
        level = item.get("risk_level") or "unknown"
        risks[level] = risks.get(level, 0) + 1
    return {"total": len(items), "by_type": counts, "by_risk": risks}


@router.get("/unified/export")
def export_unified_reports(
    report_type: str = Query("all", pattern="^(all|inspection|awr|performance|slow_sql|alert|ticket|backup|governance|ai|rescue)$"),
    limit: int = Query(100, ge=1, le=500),
    db: Session = Depends(get_db),
):
    items = _collect_unified_reports(db, report_type, limit)
    lines = [
        "# TraceMind AI 智能运维平台 统一报告快照",
        "",
        f"- 生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"- 报告范围: {report_type}",
        f"- 报告数量: {len(items)}",
        "",
    ]
    current_type = None
    for item in items:
        if item["report_type"] != current_type:
            current_type = item["report_type"]
            lines.extend(["", f"## {current_type}", ""])
        lines.extend([
            f"### {item.get('title') or item['id']}",
            f"- 实例: {item.get('instance_name') or '-'}",
            f"- 类型: {item.get('db_type') or '-'}",
            f"- 状态: {item.get('status') or '-'}",
            f"- 风险: {item.get('risk_level') or '-'}",
            f"- 时间: {item.get('create_time') or '-'}",
            f"- 摘要: {item.get('summary') or '-'}",
            "",
        ])
    stream = io.BytesIO("\n".join(lines).encode("utf-8"))
    filename = f"mesguard_unified_reports_{datetime.now().strftime('%Y%m%d_%H%M%S')}.md"
    return StreamingResponse(
        stream,
        media_type="text/markdown; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@router.post("/dbcheck")
def api_dbcheck(data: DBCheckRequest, db: Session = Depends(get_db)):
    instance = instance_service.get_by_id(db, data.instance_id)
    if not instance:
        raise HTTPException(404, "数据库实例不存在")

    try:
        result = run_dbcheck(instance)
        generated_at = datetime.now()
        result["generated_at"] = generated_at.isoformat()

        report = DBCheckReport(
            instance_id=data.instance_id,
            instance_name=instance.instance_name,
            total=result["total"],
            success=result["success"],
            failed=result["failed"],
            result_json=json.dumps(result, ensure_ascii=False, default=str),
        )
        db.add(report)
        db.commit()
        db.refresh(report)
        result["report_id"] = report.id
        return result
    except Exception as e:
        logger.error("dbcheck 执行失败 | error=%s", e, exc_info=True)
        raise HTTPException(500, f"巡检执行失败：{e}")


class InspectRequest(BaseModel):
    instance_id: int


@router.post("/inspect")
def api_unified_inspect(data: InspectRequest, db: Session = Depends(get_db)):
    """统一巡检：任意数据库类型，写入统一三表模型并返回富结果。"""
    instance = instance_service.get_by_id(db, data.instance_id)
    if not instance:
        raise HTTPException(404, "数据库实例不存在")
    try:
        return inspection_dispatch.run_through_queue(db, instance, trigger_type="manual")
    except Exception as e:
        logger.error("统一巡检失败 | instance=%s | error=%s", data.instance_id, e, exc_info=True)
        raise HTTPException(500, f"巡检执行失败：{e}")


@router.get("/inspections")
def api_list_inspections(instance_id: int = None, db_type: str = None, limit: int = 50,
                         db: Session = Depends(get_db)):
    """跨库巡检历史列表（报告中心）。"""
    return {"items": result_writer.list_tasks(db, instance_id, db_type, limit)}


@router.get("/inspections/{instance_id}/{task_id}")
def api_inspection_detail(instance_id: int, task_id: str, db: Session = Depends(get_db)):
    result = report_builder.build_report_payload(db, task_id, instance_id)
    if not result:
        raise HTTPException(404, "巡检记录不存在")
    return result


@router.get("/{task_id}/detail")
def api_inspection_structured_detail(task_id: str, instance_id: int | None = None, db: Session = Depends(get_db)):
    """结构化巡检报告详情：供前端文档式报告页使用。"""
    result = report_builder.build_report_payload(db, task_id, instance_id)
    if not result:
        raise HTTPException(404, "巡检记录不存在")
    return _build_inspection_report_detail(result)


@router.get("/{task_id}/export")
def api_inspection_report_export(
    task_id: str,
    instance_id: int | None = None,
    format: str = Query("markdown", pattern="^(markdown|html)$"),
    db: Session = Depends(get_db),
):
    """导出结构化巡检报告。PDF/DOCX 保留原 Word 接口，当前先支持 Markdown / HTML。"""
    result = report_builder.build_report_payload(db, task_id, instance_id)
    if not result:
        raise HTTPException(404, "巡检记录不存在")
    detail = _build_inspection_report_detail(result)
    if format == "html":
        content = _build_inspection_html(detail)
        suffix = "html"
        media_type = "text/html; charset=utf-8"
    else:
        content = detail["markdown"]
        suffix = "md"
        media_type = "text/markdown; charset=utf-8"

    stream = io.BytesIO(content.encode("utf-8"))
    filename = f"inspection_{task_id}.{suffix}"
    return StreamingResponse(
        stream,
        media_type=media_type,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@router.get("/inspections/{instance_id}/{task_id}/export-word")
def api_inspection_export_word(instance_id: int, task_id: str, db: Session = Depends(get_db)):
    result = report_builder.build_report_payload(db, task_id, instance_id)
    if not result:
        raise HTTPException(404, "巡检记录不存在")
    stream = _build_inspection_docx(result)
    fname = f"inspection_{instance_id}_{task_id}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


def _build_inspection_report_detail(result: dict) -> dict:
    task = dict(result.get("task") or {})
    report = result.get("report") or {}
    modules = result.get("modules") or []
    findings = result.get("findings") or []
    task_status = _inspection_status(task.get("task_status"))
    if task_status == "success" and not findings:
        task_status = "empty"
    task["task_status"] = task_status
    if task_status != "success":
        task["overall_score"] = None
        task["risk_level"] = "unknown"
    abnormal = [item for item in findings if item.get("severity") in {"critical", "warning", "high"}]
    counts = {
        "total": len(findings),
        "normal": task.get("normal_count", 0),
        "info": task.get("info_count", 0),
        "warning": task.get("warning_count", 0),
        "critical": task.get("critical_count", 0),
    }
    instance = report.get("instance") or {}
    summary = {
        **counts,
        "score": task.get("overall_score"),
        "risk_level": task.get("risk_level") or "unknown",
        "risk_text": _risk_text(task.get("risk_level")),
        "task_status": task_status,
        "failure_reason": task.get("failure_reason"),
    }
    category_summary = [
        {
            "module_key": item.get("module_key"),
            "module_name": item.get("module_name"),
            "score": item.get("score") if task_status == "success" and _inspection_status(item.get("task_status")) == "success" else None,
            "risk_level": item.get("risk_level") if task_status == "success" else "unknown",
            "summary": item.get("summary") if task_status == "success" else "未完成健康评估。",
            "counts": (item.get("evidence") or {}).get("counts") or {},
            "top_risks": item.get("top_risks") or [],
            "suggestions": item.get("suggestions") or [],
        }
        for item in modules
    ]
    ai_result = {
        "analysis_source": report.get("analysis_source") or "rule_engine",
        "root_cause_analysis": report.get("root_cause_analysis") or [],
        "business_impact": report.get("business_impact") or [],
        "recommended_priorities": report.get("recommended_priorities") or abnormal[:5],
        "ticket_description": _ticket_description(task, report, abnormal),
        "evidence_scope": [
            "巡检结构化 Finding",
            "异常项原始采集明细",
            "只读 SQL / 命令证据",
            "模块评分与风险分布",
        ],
    }
    detail = {
        "task_id": task.get("task_id"),
        "title": "TraceMind AI 数据库健康巡检报告",
        "generated_at": datetime.now().isoformat(sep=" ", timespec="seconds"),
        "instance": instance,
        "task": task,
        "summary": summary,
        "conclusion": (
            task.get("failure_reason") or f"巡检{_inspection_status_text(task_status)}，未获得完整健康评估。"
            if task_status != "success"
            else report.get("overall_conclusion") or task.get("summary")
        ),
        "risk_items": abnormal,
        "all_findings": findings,
        "category_summary": category_summary,
        "ai_result": ai_result,
        "modules": modules,
        "report": report,
        "raw_payload": result,
    }
    detail["markdown"] = _build_inspection_markdown(detail)
    return detail


def _inspection_score_text(summary: dict) -> str:
    score = summary.get("score")
    if _inspection_status(summary.get("task_status")) != "success":
        return "未生成（未完成健康评估）"
    return f"{score}/100" if score is not None else "未生成（未记录评分）"


def _inspection_status(value) -> str:
    status = str(value or "").strip().lower()
    return {
        "error": "failed", "failure": "failed", "collection_failed": "failed",
        "connection_failed": "failed", "ok": "success", "completed": "success",
        "succeeded": "success", "not_supported": "skipped", "not_applicable": "skipped",
        "not_executed": "skipped", "cancelled": "canceled",
    }.get(status, status or "unknown")


def _inspection_status_text(status: str) -> str:
    return {
        "success": "执行成功", "failed": "采集失败", "partial": "部分完成",
        "skipped": "未执行", "unknown": "未确认", "summary": "执行汇总",
        "empty": "无有效结果", "running": "运行中", "pending": "待执行",
        "plan_only": "仅生成方案", "canceled": "已取消",
    }.get(status, "未确认")


def _export_execution_status(item: dict, payload: dict) -> str:
    if payload.get("error") or payload.get("error_message") or _inspection_status(payload.get("status")) == "failed":
        return "failed"
    saved = item.get("raw_data") if isinstance(item.get("raw_data"), dict) else {}
    collected = result_writer.finding_execution_status(SimpleNamespace(
        metric_name=saved.get("metric_name"), metric_value=saved.get("metric_value"),
        check_status=saved.get("check_status"), raw_data=payload,
    ))
    statuses = [_inspection_status(value) for value in (
        item.get("execution_status"), payload.get("execution_status"), payload.get("state"),
    ) if value]
    if collected != "success":
        statuses.append(collected)
    for status in ("failed", "skipped", "partial", "unknown", "running", "pending", "plan_only", "canceled"):
        if status in statuses:
            return status
    if payload.get("evidence_kind") == "execution_summary":
        return "summary"
    if str(payload.get("evidence_source") or "").startswith("legacy_"):
        return "unknown"
    if "success" in statuses or payload.get("state") == "collected":
        return "success"
    if _evidence_sql_list(payload) and str(payload.get("status") or "").lower() in {"normal", "success", "ok"}:
        return "success"
    return "unknown"


def _export_empty_result(status: str, payload: dict) -> str:
    if status == "success":
        count = payload.get("row_count")
        if isinstance(count, (int, float)) and count > 0:
            return f"查询返回 {count} 行，未保存可展示明细。"
        if count == 0:
            return "SQL 执行成功，返回 0 行。" if _evidence_sql_list(payload) else "采集完成，返回 0 行。"
    message = payload.get("message") or payload.get("error") or payload.get("error_message")
    return "未记录可确认的查询结果。" + (str(message) if message else "")


def _export_no_risk_text(summary: dict) -> str:
    return "未发现严重或警告风险。" if summary.get("task_status") == "success" else "巡检未完成健康评估，不能据此认定无风险。"


def _build_inspection_markdown(detail: dict) -> str:
    task = detail.get("task") or {}
    instance = detail.get("instance") or {}
    report = detail.get("report") or {}
    summary = detail.get("summary") or {}
    risk_items = detail.get("risk_items") or []
    categories = detail.get("category_summary") or []
    all_findings = detail.get("all_findings") or []
    ai_result = detail.get("ai_result") or {}
    instance_address = f"{instance.get('host') or '-'}:{instance.get('port') or '-'}"
    lines = [
        "# TraceMind AI 数据库健康巡检报告",
        "",
        "> TraceMind AI 智能运维平台 · 数据库健康检查归档",
        "",
        "## 报告信息",
        "",
        "| 项目 | 内容 | 项目 | 内容 |",
        "| --- | --- | --- | --- |",
        f"| 实例名称 | {_markdown_cell(instance.get('name'))} | 数据库类型 | {_markdown_cell(instance.get('db_type') or task.get('db_type'))} |",
        f"| 连接地址 | {_markdown_cell(instance_address)} | 任务号 | {_markdown_cell(task.get('task_id'))} |",
        f"| 巡检时间 | {_markdown_cell(report.get('inspection_time') or task.get('finished_at') or task.get('create_time'))} | 生成时间 | {_markdown_cell(detail.get('generated_at'))} |",
        f"| 巡检模板 | {_markdown_cell(task.get('template_name') or task.get('template_key') or '未记录（历史任务）')} | 账号能力 | {_markdown_cell(task.get('account_capability_label') or ('只读降级采集' if task.get('read_only') else '账号能力未评估'))} |",
        "",
        "## 执行摘要",
        "",
        "| 健康分 | 风险等级 | 总检查项 | 正常 | 提示 | 警告 | 严重 |",
        "| ---: | --- | ---: | ---: | ---: | ---: | ---: |",
        f"| {_inspection_score_text(summary)} | {_risk_text(summary.get('risk_level'))} | {summary.get('total', 0)} | {summary.get('normal', 0)} | {summary.get('info', 0)} | {summary.get('warning', 0)} | {summary.get('critical', 0)} |",
        "",
        "> 规则健康分用于风险排序：严重项扣 25 分、警告项扣 10 分、提示项扣 2 分；它不等同于数据库可用性 SLA。",
        "",
        "### 综合结论",
        "",
        _md(detail.get("conclusion") or "-"),
        "",
        "## 风险与处置建议",
        "",
    ]
    if risk_items:
        for index, item in enumerate(risk_items, 1):
            suggestions = item.get("suggestions") or []
            lines.extend([
                f"### {index}. [{_risk_text(item.get('severity'))}] {_md(item.get('check_item'))}",
                f"- 发现: {_md(item.get('finding'))}",
                f"- 建议: {_md('；'.join(suggestions) if suggestions else '-')}",
                f"- 处置状态: {_md((item.get('operation_plan') or {}).get('status') or 'plan_only')}",
                "",
            ])
    else:
        lines.extend([_export_no_risk_text(summary), ""])

    lines.extend([
        "## 模块评估",
        "",
        "| 模块 | 得分 | 风险 | 严重 | 警告 | 提示 | 正常 | 结论 |",
        "| --- | ---: | --- | ---: | ---: | ---: | ---: | --- |",
    ])
    for item in categories:
        counts = item.get("counts") or {}
        lines.append(
            "| "
            + " | ".join(
                [
                    _markdown_cell(item.get("module_name")),
                    _markdown_cell(item.get("score")),
                    _markdown_cell(_risk_text(item.get("risk_level"))),
                    _markdown_cell(counts.get("critical", 0)),
                    _markdown_cell(counts.get("warning", 0)),
                    _markdown_cell(counts.get("info", 0)),
                    _markdown_cell(counts.get("normal", 0)),
                    _markdown_cell(item.get("summary")),
                ]
            )
            + " |"
        )
    if not categories:
        lines.append("| 暂无模块评分 | - | - | 0 | 0 | 0 | 0 | - |")

    lines.extend(["", "## SQL 与查询结果", ""])
    evidence_count = 0
    for index, item in enumerate(all_findings, 1):
        payload = _finding_evidence(item)
        sql_list = _evidence_sql_list(payload)
        columns, rows = _evidence_table(payload)
        evidence_count += 1
        status = _export_execution_status(item, payload)
        row_count = payload.get("row_count") if payload.get("row_count") is not None else "未记录"
        lines.extend([
            f"### {index}. {_md(item.get('check_item') or payload.get('title') or '巡检项')}",
            f"- 模块: {_md(item.get('module_key'))}",
            f"- 执行状态: {_inspection_status_text(status)}",
            f"- 返回行数: {_md(row_count)}",
            f"- 脚本行号: {_md(payload.get('line_number'))}",
            f"- 检查结论: {_md(item.get('finding'))}",
        ])
        if payload.get("message"):
            lines.append(f"- 执行说明: {_md(payload.get('message'))}")
        if payload.get("truncated"):
            lines.append(f"- 结果截断: 是，报告保留前 {_md(payload.get('max_rows') or len(rows))} 行")
        if sql_list:
            lines.extend(["", "**执行 SQL / 采集命令**", ""])
            for sql in sql_list:
                lines.extend(["```sql", str(sql), "```", ""])
        lines.append("**查询结果**")
        lines.append("")
        if rows and columns:
            lines.extend(_markdown_result_table(columns, rows))
        else:
            lines.append(f"_{_md(_export_empty_result(status, payload))}_")
        lines.append("")
    if evidence_count == 0:
        lines.extend(["本次报告没有可展示的 SQL 或查询结果。", ""])

    lines.extend(["## 综合分析（规则引擎）", ""])
    root_causes = ai_result.get("root_cause_analysis") or []
    impacts = ai_result.get("business_impact") or []
    priorities = ai_result.get("recommended_priorities") or []
    if root_causes or impacts or priorities:
        lines.append("### 风险成因")
        lines.extend([f"- {_md(item)}" for item in root_causes] or ["- 暂无风险成因分析"])
        lines.extend(["", "### 影响范围"])
        lines.extend([f"- {_md(item)}" for item in impacts] or ["- 暂无业务影响"])
        lines.extend(["", "### 优先处置建议"])
        lines.extend([f"- {_md(item.get('check_item') or item.get('finding'))}" for item in priorities] or ["- 暂无优先处置项"])
    else:
        evidence_scope = "、".join(str(value) for value in ai_result.get("evidence_scope") or [])
        lines.append(
            _export_no_risk_text(summary)
            + (f"已纳入分析的证据范围：{evidence_scope}。" if evidence_scope else "")
        )
    lines.extend([
        "",
        "## 工单描述",
        "",
        _md(ai_result.get("ticket_description") or "-"),
        "",
        "---",
        "",
        f"生成系统：TraceMind AI · 生成时间：{_md(detail.get('generated_at'))}",
        "",
    ])
    return "\n".join(lines)


def _build_inspection_html(detail: dict) -> str:
    task = detail.get("task") or {}
    instance = detail.get("instance") or {}
    report = detail.get("report") or {}
    summary = detail.get("summary") or {}
    risks = detail.get("risk_items") or []
    categories = detail.get("category_summary") or []
    findings = detail.get("all_findings") or []
    ai_result = detail.get("ai_result") or {}
    title = _html_text(detail.get("title") or "TraceMind AI 数据库健康巡检报告")
    db_type = _html_text(instance.get("db_type") or task.get("db_type"))
    instance_name = _html_text(instance.get("name") or task.get("instance_id"))
    address = _html_text(f"{instance.get('host') or '-'}:{instance.get('port') or '-'}")
    risk_key = _html_risk_key(summary.get("risk_level"))
    metadata = [
        ("实例名称", instance_name),
        ("连接地址", address),
        ("数据库类型", db_type.upper()),
        ("任务号", _html_text(task.get("task_id"))),
        ("巡检时间", _html_text(report.get("inspection_time") or task.get("finished_at") or task.get("create_time"))),
        ("生成时间", _html_text(detail.get("generated_at"))),
        ("巡检模板", _html_text(task.get("template_name") or task.get("template_key") or "未记录（历史任务）")),
        (
            "账号能力",
            _html_text(
                task.get("account_capability_label")
                or ("只读降级采集" if task.get("read_only") else "账号能力未评估")
            ),
        ),
    ]
    metadata_html = "".join(
        f'<div class="meta-item"><span>{label}</span><strong>{value}</strong></div>'
        for label, value in metadata
    )
    metrics = [
        ("健康分", _inspection_score_text(summary), risk_key),
        ("总检查项", summary.get("total", 0), ""),
        ("正常", summary.get("normal", 0), "normal"),
        ("提示", summary.get("info", 0), "info"),
        ("警告", summary.get("warning", 0), "warning"),
        ("严重", summary.get("critical", 0), "critical"),
    ]
    metrics_html = "".join(
        f'<div class="metric {css_class}"><span>{label}</span><strong>{_html_text(value)}</strong></div>'
        for label, value, css_class in metrics
    )
    if risks:
        risks_html = "".join(
            f"""
            <article class="risk-card {_html_risk_key(item.get('severity'))}">
              <div class="risk-title">
                <span class="status">{_html_text(_risk_text(item.get('severity')))}</span>
                <h3>{index}. {_html_text(item.get('check_item') or '巡检项')}</h3>
              </div>
              <p>{_html_text(item.get('finding'))}</p>
              <div class="recommendation"><b>处置建议</b>{_html_text('；'.join(item.get('suggestions') or []) or '-')}</div>
            </article>
            """
            for index, item in enumerate(risks, 1)
        )
    else:
        risks_html = f'<div class="empty-state">{_export_no_risk_text(summary)}</div>'
    category_rows = []
    for item in categories:
        counts = item.get("counts") or {}
        category_rows.append(
            [
                item.get("module_name"),
                item.get("score"),
                _risk_text(item.get("risk_level")),
                counts.get("critical", 0),
                counts.get("warning", 0),
                counts.get("info", 0),
                counts.get("normal", 0),
                item.get("summary"),
            ]
        )
    categories_html = _html_table(
        ["模块", "得分", "风险", "严重", "警告", "提示", "正常", "结论"],
        category_rows,
        "module-table",
    )
    evidence_sections = []
    for index, item in enumerate(findings, 1):
        payload = _finding_evidence(item)
        sql_list = _evidence_sql_list(payload)
        columns, rows = _evidence_table(payload)
        status = _export_execution_status(item, payload)
        row_count = payload.get("row_count") if payload.get("row_count") is not None else "未记录"
        sql_html = "".join(
            f"<pre><code>{_html_text(sql)}</code></pre>"
            for sql in sql_list
        ) or '<div class="empty-inline">未记录执行 SQL。</div>'
        result_html = (
            _html_table(columns, [[row.get(column) for column in columns] for row in rows], "result-table")
            if rows and columns
            else f'<div class="empty-state">{_html_text(_export_empty_result(status, payload))}</div>'
        )
        evidence_sections.append(
            f"""
            <article class="evidence-card">
              <div class="evidence-head">
                <div><span class="sequence">{index:02d}</span><h3>{_html_text(item.get('check_item') or payload.get('title') or '巡检项')}</h3></div>
                <span class="status {_html_risk_key(status)}">{_inspection_status_text(status)}</span>
              </div>
              <dl class="evidence-meta">
                <div><dt>模块</dt><dd>{_html_text(item.get('module_key'))}</dd></div>
                <div><dt>返回行数</dt><dd>{_html_text(row_count)}</dd></div>
                <div><dt>脚本行号</dt><dd>{_html_text(payload.get('line_number'))}</dd></div>
                <div><dt>结果截断</dt><dd>{'是' if payload.get('truncated') else '否'}</dd></div>
              </dl>
              {_html_callout(item.get('finding'))}
              {_html_callout(payload.get('summary') or payload.get('message'))}
              <h4>执行 SQL / 采集命令</h4>
              {sql_html}
              <h4>查询结果</h4>
              <div class="table-scroll">{result_html}</div>
            </article>
            """
        )
    evidence_html = "".join(evidence_sections) or '<div class="empty-state">本次报告没有可展示的 SQL 或查询结果。</div>'
    ai_sections = [
        ("风险成因", ai_result.get("root_cause_analysis") or []),
        ("影响范围", ai_result.get("business_impact") or []),
        (
            "优先处置建议",
            [
                item.get("check_item") or item.get("finding")
                for item in ai_result.get("recommended_priorities") or []
            ],
        ),
    ]
    if any(items for _, items in ai_sections):
        ai_html = "".join(
            f'<div class="ai-block"><h3>{label}</h3>{_html_list(items)}</div>'
            for label, items in ai_sections
        )
    else:
        evidence_scope = "、".join(str(value) for value in ai_result.get("evidence_scope") or [])
        ai_html = (
            '<div class="summary ai-summary">'
            + _export_no_risk_text(summary)
            + (f" 已纳入分析的证据范围：{_html_text(evidence_scope)}。" if evidence_scope else "")
            + "</div>"
        )
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>{title}</title>
  <style>
    :root {{ color-scheme: light; --ink: #1d1d1f; --muted: #6e6e73; --line: #e2e3e7; --green: #16794c; --blue: #1769aa; --orange: #b85c00; --red: #c3272b; }}
    * {{ box-sizing: border-box; }}
    body {{ margin: 0; background: #f5f5f7; color: var(--ink); font: 14px/1.65 "Microsoft YaHei", "PingFang SC", Arial, sans-serif; }}
    main {{ width: min(1180px, calc(100% - 40px)); margin: 32px auto; padding: 40px 44px; background: #fff; border: 1px solid var(--line); border-radius: 10px; box-shadow: 0 12px 36px rgb(0 0 0 / 7%); }}
    header {{ display: flex; gap: 24px; align-items: end; justify-content: space-between; padding-bottom: 24px; border-bottom: 2px solid #143d33; }}
    .eyebrow {{ color: var(--green); font-size: 12px; font-weight: 700; letter-spacing: .08em; }}
    h1 {{ margin: 4px 0 2px; font-size: 28px; line-height: 1.3; letter-spacing: 0; }}
    h2 {{ margin: 34px 0 14px; font-size: 19px; line-height: 1.35; letter-spacing: 0; }}
    h3, h4 {{ letter-spacing: 0; }}
    .report-id {{ color: var(--muted); font-size: 12px; text-align: right; }}
    .status {{ display: inline-block; padding: 2px 8px; border: 1px solid #d8d9dd; border-radius: 4px; background: #f7f7f8; font-size: 12px; white-space: nowrap; }}
    .status.normal, .metric.normal strong {{ color: var(--green); }}
    .status.info, .metric.info strong {{ color: var(--blue); }}
    .status.warning, .metric.warning strong {{ color: var(--orange); }}
    .status.critical, .metric.critical strong {{ color: var(--red); }}
    .meta-grid {{ display: grid; grid-template-columns: repeat(2, minmax(0, 1fr)); border: 1px solid var(--line); border-radius: 6px; overflow: hidden; margin-top: 22px; }}
    .meta-item {{ display: grid; grid-template-columns: 108px minmax(0, 1fr); min-height: 44px; border-bottom: 1px solid var(--line); }}
    .meta-item:nth-child(odd) {{ border-right: 1px solid var(--line); }}
    .meta-item:nth-last-child(-n + 2) {{ border-bottom: 0; }}
    .meta-item span, .meta-item strong {{ padding: 11px 13px; }}
    .meta-item span {{ color: var(--muted); background: #fafafa; font-weight: 400; }}
    .meta-item strong {{ overflow-wrap: anywhere; font-weight: 500; }}
    .metrics {{ display: grid; grid-template-columns: repeat(6, 1fr); margin-top: 18px; border: 1px solid var(--line); border-radius: 6px; overflow: hidden; }}
    .metric {{ min-width: 0; padding: 13px 14px; border-right: 1px solid var(--line); }}
    .metric:last-child {{ border-right: 0; }}
    .metric span {{ display: block; color: var(--muted); font-size: 12px; }}
    .metric strong {{ display: block; margin-top: 2px; font-size: 23px; line-height: 1.3; }}
    .summary {{ padding: 15px 18px; border-left: 4px solid #143d33; background: #f3f8f6; }}
    .risk-list {{ display: grid; gap: 10px; }}
    .risk-card {{ padding: 15px 17px; border: 1px solid var(--line); border-left-width: 4px; border-radius: 6px; break-inside: avoid; }}
    .risk-card.warning {{ border-left-color: var(--orange); }}
    .risk-card.critical {{ border-left-color: var(--red); }}
    .risk-title, .evidence-head, .risk-title > div, .evidence-head > div {{ display: flex; gap: 10px; align-items: center; justify-content: space-between; }}
    .risk-title h3, .evidence-head h3 {{ flex: 1; margin: 0; font-size: 15px; }}
    .risk-card p {{ margin: 9px 0; }}
    .recommendation {{ display: grid; grid-template-columns: 76px 1fr; gap: 8px; color: #444; }}
    table {{ width: 100%; border-collapse: collapse; table-layout: auto; }}
    th, td {{ padding: 8px 10px; border: 1px solid var(--line); text-align: left; vertical-align: top; overflow-wrap: anywhere; }}
    th {{ color: #3a3a3c; background: #f1f3f5; font-size: 12px; font-weight: 600; }}
    td {{ font-size: 12px; }}
    .table-scroll {{ overflow-x: auto; }}
    .result-table {{ min-width: 680px; }}
    .evidence-list {{ display: grid; gap: 16px; }}
    .evidence-card {{ padding: 18px; border: 1px solid var(--line); border-radius: 6px; break-inside: avoid; }}
    .sequence {{ color: var(--muted); font: 600 12px/1.2 ui-monospace, monospace; }}
    .evidence-meta {{ display: grid; grid-template-columns: repeat(4, 1fr); margin: 12px 0; border: 1px solid var(--line); }}
    .evidence-meta div {{ padding: 8px 10px; border-right: 1px solid var(--line); }}
    .evidence-meta div:last-child {{ border-right: 0; }}
    dt {{ color: var(--muted); font-size: 11px; }}
    dd {{ margin: 1px 0 0; font-size: 12px; font-weight: 600; }}
    .callout {{ margin: 10px 0; padding: 10px 12px; color: #4c3a17; background: #fff8e8; border: 1px solid #ecd8aa; }}
    pre {{ margin: 8px 0 14px; padding: 14px 16px; overflow: auto; color: #d9fbe8; background: #111827; border-radius: 5px; white-space: pre-wrap; word-break: break-word; font: 12px/1.55 "SFMono-Regular", Consolas, monospace; }}
    .empty-state, .empty-inline {{ padding: 18px; color: var(--muted); background: #fafafa; border: 1px dashed #d9dade; text-align: center; }}
    .empty-inline {{ padding: 10px; text-align: left; }}
    .ai-grid {{ display: grid; grid-template-columns: repeat(3, 1fr); gap: 12px; }}
    .ai-summary {{ grid-column: 1 / -1; }}
    .ai-block {{ padding: 15px; border: 1px solid var(--line); border-radius: 6px; }}
    .ai-block h3 {{ margin: 0 0 8px; font-size: 14px; }}
    .ai-block ul {{ margin: 0; padding-left: 20px; }}
    footer {{ margin-top: 34px; padding-top: 14px; color: var(--muted); border-top: 1px solid var(--line); font-size: 11px; text-align: right; }}
    @media (max-width: 800px) {{
      main {{ width: 100%; margin: 0; padding: 24px 18px; border: 0; border-radius: 0; }}
      header {{ align-items: flex-start; flex-direction: column; }}
      .report-id {{ text-align: left; }}
      .meta-grid, .ai-grid {{ grid-template-columns: 1fr; }}
      .meta-item:nth-child(odd) {{ border-right: 0; }}
      .meta-item:nth-last-child(2) {{ border-bottom: 1px solid var(--line); }}
      .metrics {{ grid-template-columns: repeat(2, 1fr); }}
      .metric {{ border-bottom: 1px solid var(--line); }}
      .evidence-meta {{ grid-template-columns: repeat(2, 1fr); }}
    }}
    @media print {{
      @page {{ size: A4 landscape; margin: 14mm; }}
      body {{ background: #fff; }}
      main {{ width: 100%; margin: 0; padding: 0; border: 0; box-shadow: none; }}
      .evidence-card, .risk-card, table {{ break-inside: avoid; }}
      pre {{ color: #111; background: #f3f4f6; border: 1px solid #d8dadd; }}
    }}
  </style>
</head>
<body>
  <main>
    <header>
      <div>
        <div class="eyebrow">DATABASE HEALTH CHECK</div>
        <h1>{title}</h1>
        <span class="status {risk_key}">{_html_text(_risk_text(summary.get("risk_level")))}</span>
      </div>
      <div class="report-id">任务号<br><strong>{_html_text(task.get("task_id"))}</strong></div>
    </header>
    <section class="meta-grid">{metadata_html}</section>
    <section class="metrics">{metrics_html}</section>
    <h2>执行摘要</h2>
    <div class="summary">{_html_text(detail.get("conclusion") or "-")}</div>
    <h2>风险与处置建议</h2>
    <section class="risk-list">{risks_html}</section>
    <h2>模块评估</h2>
    <div class="table-scroll">{categories_html}</div>
    <h2>SQL 与查询结果</h2>
    <section class="evidence-list">{evidence_html}</section>
    <h2>综合分析（规则引擎）</h2>
    <section class="ai-grid">{ai_html}</section>
    <h2>工单描述</h2>
    <div class="summary">{_html_text(ai_result.get("ticket_description") or "-")}</div>
    <footer>TraceMind AI 智能运维平台 · 生成时间 {_html_text(detail.get("generated_at"))}</footer>
  </main>
</body>
</html>"""


def _html_text(value) -> str:
    if value is None or value == "":
        return "-"
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    return html.escape(str(value), quote=True)


def _html_risk_key(level) -> str:
    value = str(level or "normal").lower()
    if value in {"critical", "high", "disaster", "error", "failed"}:
        return "critical"
    if value in {"warning", "average", "skipped"}:
        return "warning"
    if value in {"info", "low", "notice"}:
        return "info"
    return "normal"


def _html_table(headers: list, rows: list, class_name: str = "") -> str:
    if not rows:
        return '<div class="empty-state">暂无数据。</div>'
    head = "".join(f"<th>{_html_text(item)}</th>" for item in headers)
    body = "".join(
        "<tr>" + "".join(f"<td>{_html_text(value)}</td>" for value in row) + "</tr>"
        for row in rows
    )
    return f'<table class="{html.escape(class_name, quote=True)}"><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>'


def _html_callout(value) -> str:
    if value in (None, ""):
        return ""
    return f'<div class="callout">{_html_text(value)}</div>'


def _html_list(items: list) -> str:
    values = [item for item in items if item not in (None, "")]
    if not values:
        return '<div class="empty-inline">暂无数据。</div>'
    return "<ul>" + "".join(f"<li>{_html_text(item)}</li>" for item in values) + "</ul>"


def _ticket_description(task: dict, report: dict, abnormal: list[dict]) -> str:
    instance = report.get("instance") or {}
    score_text = _inspection_score_text({"score": task.get("overall_score"), "task_status": task.get("task_status")})
    lines = [
        f"实例：{instance.get('name') or task.get('instance_id')}（{instance.get('db_type') or task.get('db_type')}）",
        f"任务号：{task.get('task_id')}",
        f"健康分：{score_text}，风险等级：{_risk_text(task.get('risk_level'))}",
        f"异常概览：Critical {task.get('critical_count', 0)}，Warning {task.get('warning_count', 0)}。",
    ]
    for item in abnormal[:5]:
        lines.append(f"- [{_risk_text(item.get('severity'))}] {item.get('check_item')}: {item.get('finding')}")
    lines.append("处置要求：先复核证据链，所有变更保持 plan_only，审批通过后在维护窗口人工执行。")
    return "\n".join(lines)


def _risk_text(level) -> str:
    return {
        "critical": "严重",
        "high": "高风险",
        "warning": "警告",
        "info": "提示",
        "normal": "健康",
    }.get(str(level or "").lower(), str(level or "-"))


def _md(value) -> str:
    if value is None:
        return "-"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value).replace("\n", " ").strip() or "-"


def _finding_evidence(item: dict) -> dict:
    evidence = item.get("evidence")
    if isinstance(evidence, dict):
        return evidence
    raw_data = item.get("raw_data")
    if isinstance(raw_data, dict):
        nested = raw_data.get("raw_data")
        if isinstance(nested, dict):
            return nested
        return raw_data
    return {}


def _evidence_sql_list(payload: dict) -> list[str]:
    sql = payload.get("sql")
    if isinstance(sql, list):
        return [str(item) for item in sql if item not in (None, "")]
    return [str(sql)] if sql not in (None, "") else []


def _evidence_table(payload: dict) -> tuple[list[str], list[dict]]:
    source_rows = payload.get("rows")
    if not isinstance(source_rows, list):
        source_rows = payload.get("items")
    if not isinstance(source_rows, list):
        source_rows = []
    rows = [
        row if isinstance(row, dict) else {"value": row}
        for row in source_rows[:200]
    ]
    columns: list[str] = []
    for column in payload.get("columns") or []:
        if isinstance(column, dict):
            key = column.get("name") or column.get("key") or column.get("title")
        else:
            key = column
        key = str(key or "").strip()
        if key and key not in columns:
            columns.append(key)
    for row in rows[:20]:
        for key in row:
            key = str(key)
            if key not in columns:
                columns.append(key)
    if rows and not columns:
        columns = ["value"]
    return columns, rows


def _markdown_cell(value) -> str:
    if value is None:
        return "-"
    if isinstance(value, (dict, list)):
        text = json.dumps(value, ensure_ascii=False, default=str)
    else:
        text = str(value)
    return text.replace("\\", "\\\\").replace("|", "\\|").replace("\r\n", "<br>").replace("\n", "<br>") or "-"


def _markdown_result_table(columns: list[str], rows: list[dict]) -> list[str]:
    lines = [
        "| " + " | ".join(_markdown_cell(column) for column in columns) + " |",
        "| " + " | ".join("---" for _ in columns) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(_markdown_cell(row.get(column)) for column in columns) + " |")
    return lines


def _build_inspection_docx(result: dict) -> io.BytesIO:
    import sys

    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    detail = _build_inspection_report_detail(result)
    task = detail.get("task") or {}
    instance = detail.get("instance") or {}
    report = detail.get("report") or {}
    summary = detail.get("summary") or {}
    modules = detail.get("category_summary") or []
    findings = detail.get("all_findings") or []
    risks = detail.get("risk_items") or []
    ai_result = detail.get("ai_result") or {}
    ink = "1D1D1F"
    muted = "6E6E73"
    green = "16794C"
    dark_green = "143D33"
    blue = "1769AA"
    orange = "B85C00"
    red = "C3272B"
    line = "D9DCE1"
    light_fill = "F3F5F7"

    body_font = "Arial Unicode MS" if sys.platform == "darwin" else "Microsoft YaHei"

    def set_run_font(run, size=None, bold=None, color=None, name=body_font):
        run.font.name = name
        # SQL's Latin monospace font must not replace the Chinese fallback.
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), body_font)
        if size is not None:
            run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        return run

    def shade_cell(cell, fill):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), fill)

    def set_cell_margin(cell, top=70, start=100, bottom=70, end=100):
        tc = cell._tc
        tc_pr = tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in("w:tcMar")
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
            node = tc_mar.find(qn(f"w:{margin_name}"))
            if node is None:
                node = OxmlElement(f"w:{margin_name}")
                tc_mar.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

    def set_cell_text(cell, value, *, bold=False, color=ink, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT):
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        run = paragraph.add_run(_md(value))
        set_run_font(run, size=size, bold=bold, color=color)

    def set_repeat_table_header(row):
        tr_pr = row._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True

    def prevent_row_split(row):
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)

    def set_table_widths(table, widths_cm):
        table.autofit = False
        for row in table.rows:
            for cell, width in zip(row.cells, widths_cm):
                cell.width = Cm(width)
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_width = tc_pr.first_child_found_in("w:tcW")
                if tc_width is None:
                    tc_width = OxmlElement("w:tcW")
                    tc_pr.append(tc_width)
                tc_width.set(qn("w:w"), str(round(width * 566.929)))
                tc_width.set(qn("w:type"), "dxa")
        grid = table._tbl.tblGrid
        for grid_col, width in zip(grid.gridCol_lst, widths_cm):
            grid_col.set(qn("w:w"), str(round(width * 566.929)))

    def add_table(headers, rows, widths_cm, *, body_size=8.5, header_fill="E8EEF5"):
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for cell, header in zip(table.rows[0].cells, headers):
            set_cell_text(cell, header, bold=True, color=dark_green, size=8)
            shade_cell(cell, header_fill)
            set_cell_margin(cell)
        set_repeat_table_header(table.rows[0])
        prevent_row_split(table.rows[0])
        for row_values in rows:
            row = table.add_row()
            prevent_row_split(row)
            for cell, value in zip(row.cells, row_values):
                set_cell_text(cell, value, size=body_size)
                set_cell_margin(cell)
        set_table_widths(table, widths_cm)
        return table

    def result_table_widths(headers, rows):
        if not headers:
            return []
        minimum = 1.75 if len(headers) >= 7 else 2.25
        available = max(0.0, 25.2 - minimum * len(headers))
        weights = []
        for index, header in enumerate(headers):
            content_lengths = [
                len(_md(row[index]))
                for row in rows[:30]
                if index < len(row)
            ]
            observed = max([len(str(header)), *content_lengths], default=len(str(header)))
            weights.append(min(8.0, max(1.0, observed ** 0.5 / 2)))
        total_weight = sum(weights) or 1
        return [minimum + available * weight / total_weight for weight in weights]

    def add_section_heading(text, level=1):
        paragraph = doc.add_heading(text, level=level)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
        return paragraph

    def add_note(label, value, *, fill="F3F8F6", accent=dark_green):
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        set_cell_text(table.cell(0, 0), label, bold=True, color=accent, size=9)
        set_cell_text(table.cell(0, 1), value, color=ink, size=9)
        shade_cell(table.cell(0, 0), fill)
        shade_cell(table.cell(0, 1), fill)
        set_cell_margin(table.cell(0, 0), top=110, bottom=110)
        set_cell_margin(table.cell(0, 1), top=110, bottom=110)
        set_table_widths(table, [2.7, 22.5])
        prevent_row_split(table.rows[0])
        return table

    def add_spacer(points=4):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(points)
        paragraph.paragraph_format.line_spacing = 0.2

    doc = Document()
    doc.core_properties.title = "TraceMind AI 数据库健康巡检报告"
    doc.core_properties.subject = f"{instance.get('name') or task.get('instance_id')} 数据库健康巡检"
    doc.core_properties.author = "TraceMind AI"
    doc.core_properties.keywords = "数据库, 巡检, 运维, TraceMind AI"

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = body_font
    normal.font.size = Pt(9)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), body_font)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size, color in (
        ("Title", 24, dark_green),
        ("Heading 1", 16, dark_green),
        ("Heading 2", 11.5, "2F5F54"),
        ("Heading 3", 10, ink),
    ):
        style = doc.styles[style_name]
        style.font.name = body_font
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), body_font)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(2)
    set_run_font(header.add_run("TraceMind AI  |  数据库健康巡检"), size=8, bold=True, color=dark_green)
    header_border = OxmlElement("w:pBdr")
    border_bottom = OxmlElement("w:bottom")
    border_bottom.set(qn("w:val"), "single")
    border_bottom.set(qn("w:sz"), "6")
    border_bottom.set(qn("w:space"), "4")
    border_bottom.set(qn("w:color"), "B7C9C2")
    header_border.append(border_bottom)
    header._p.get_or_add_pPr().append(header_border)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run("内部运维资料  |  TraceMind AI  |  第 "), size=8, color=muted)
    page_run = footer.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    page_run._r.append(begin)
    page_run._r.append(instruction)
    page_run._r.append(end)
    set_run_font(page_run, size=8, color=muted)
    set_run_font(footer.add_run(" 页"), size=8, color=muted)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(2)
    set_run_font(kicker.add_run("DATABASE HEALTH CHECK"), size=8.5, bold=True, color=green)
    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(3)
    set_run_font(title.add_run("TraceMind AI 数据库健康巡检报告"), size=24, bold=True, color=dark_green)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    address = f"{instance.get('host') or '-'}:{instance.get('port') or '-'}"
    set_run_font(
        subtitle.add_run(
            f"{instance.get('name') or task.get('instance_id') or '-'}  |  "
            f"{str(instance.get('db_type') or task.get('db_type') or '-').upper()}  |  {address}"
        ),
        size=10,
        color=muted,
    )

    inspection_time = report.get("inspection_time") or task.get("finished_at") or task.get("create_time")
    account_capability = (
        task.get("account_capability_label")
        or ("只读降级采集" if task.get("read_only") else "账号能力未评估")
    )
    metadata_rows = [
        ["实例名称", instance.get("name") or task.get("instance_id"), "连接地址", address],
        ["数据库类型", str(instance.get("db_type") or task.get("db_type") or "-").upper(), "任务号", task.get("task_id")],
        ["巡检时间", inspection_time, "生成时间", detail.get("generated_at")],
        [
            "巡检模板",
            task.get("template_name") or task.get("template_key") or "未记录（历史任务）",
            "账号能力",
            account_capability,
        ],
    ]
    meta_table = doc.add_table(rows=0, cols=4)
    meta_table.style = "Table Grid"
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for values in metadata_rows:
        row = meta_table.add_row()
        prevent_row_split(row)
        for index, (cell, value) in enumerate(zip(row.cells, values)):
            is_label = index in {0, 2}
            set_cell_text(
                cell,
                value,
                bold=is_label,
                color=muted if is_label else ink,
                size=8.5,
            )
            if is_label:
                shade_cell(cell, light_fill)
            set_cell_margin(cell, top=90, bottom=90)
    set_table_widths(meta_table, [2.5, 10.1, 2.5, 10.1])
    add_spacer(5)

    metric_labels = ["健康分", "总检查项", "正常", "提示", "警告", "严重"]
    metric_values = [
        _inspection_score_text(summary),
        summary.get("total", 0),
        summary.get("normal", 0),
        summary.get("info", 0),
        summary.get("warning", 0),
        summary.get("critical", 0),
    ]
    metric_colors = [dark_green, ink, green, blue, orange, red]
    metric_table = doc.add_table(rows=2, cols=6)
    metric_table.style = "Table Grid"
    metric_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, label in enumerate(metric_labels):
        set_cell_text(metric_table.cell(0, index), label, bold=True, color=muted, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(metric_table.cell(0, index), light_fill)
        set_cell_text(
            metric_table.cell(1, index),
            metric_values[index],
            bold=True,
            color=metric_colors[index],
            size=9 if index == 0 and summary.get("score") is None else 15,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_margin(metric_table.cell(0, index), top=70, bottom=55)
        set_cell_margin(metric_table.cell(1, index), top=85, bottom=85)
    set_table_widths(metric_table, [4.2] * 6)
    prevent_row_split(metric_table.rows[0])
    prevent_row_split(metric_table.rows[1])

    add_section_heading("执行摘要")
    add_note("综合结论", detail.get("conclusion") or "-", fill="F1F7F4", accent=dark_green)

    add_section_heading("风险与处置建议")
    if risks:
        for index, item in enumerate(risks, 1):
            risk_key = _html_risk_key(item.get("severity"))
            accent = red if risk_key == "critical" else orange if risk_key == "warning" else blue
            add_note(
                f"{index}. {_risk_text(item.get('severity'))}",
                f"{item.get('check_item') or '巡检项'}：{item.get('finding') or '-'}",
                fill="FFF7E8" if risk_key == "warning" else "FFF1F0",
                accent=accent,
            )
            suggestions = item.get("suggestions") or []
            if suggestions:
                paragraph = doc.add_paragraph(style="List Bullet")
                set_run_font(
                    paragraph.add_run("处置建议：" + "；".join(str(value) for value in suggestions)),
                    size=8.5,
                    color=ink,
                )
    else:
        add_note("风险状态", _export_no_risk_text(summary), fill="F1F8F4", accent=green)

    add_section_heading("模块评估")
    if modules:
        module_rows = []
        for module in modules:
            counts = module.get("counts") or {}
            module_rows.append(
                [
                    module.get("module_name"),
                    module.get("score"),
                    _risk_text(module.get("risk_level")),
                    f"严重 {counts.get('critical', 0)} / 警告 {counts.get('warning', 0)} / 提示 {counts.get('info', 0)} / 正常 {counts.get('normal', 0)}",
                    module.get("summary"),
                ]
            )
        add_table(
            ["模块", "得分", "风险", "异常分布", "结论"],
            module_rows,
            [5.0, 1.8, 2.2, 7.0, 9.2],
            body_size=8,
        )
    else:
        add_note("模块评分", "暂无模块评分数据。", fill=light_fill, accent=muted)

    sql_heading = add_section_heading(f"SQL 与查询结果（{len(findings)} 项）")
    sql_heading.paragraph_format.page_break_before = True
    intro = doc.add_paragraph()
    set_run_font(
        intro.add_run("以下内容按检查项保留执行状态、采集 SQL、返回行数与查询结果，便于复核和审计。"),
        size=8.5,
        color=muted,
    )
    for index, finding in enumerate(findings, 1):
        item_heading = add_section_heading(
            f"{index}. {finding.get('check_item') or '巡检项'}",
            level=2,
        )
        severity = finding.get("severity")
        risk_key = _html_risk_key(severity)
        heading_color = red if risk_key == "critical" else orange if risk_key == "warning" else green
        for run in item_heading.runs:
            set_run_font(run, size=11.5, bold=True, color=heading_color)
        payload = _finding_evidence(finding)
        sql_list = _evidence_sql_list(payload)
        columns, rows = _evidence_table(payload)
        status = _export_execution_status(finding, payload)
        row_count = payload.get("row_count") if payload.get("row_count") is not None else "未记录"
        add_table(
            ["模块", "执行状态", "返回行数", "脚本行号"],
            [[finding.get("module_key"), _inspection_status_text(status), row_count, payload.get("line_number")]],
            [9.4, 5.4, 5.2, 5.2],
            body_size=8,
        )
        if finding.get("finding"):
            finding_paragraph = doc.add_paragraph()
            finding_paragraph.paragraph_format.keep_with_next = bool(sql_list or rows)
            set_run_font(finding_paragraph.add_run("检查结论："), size=8.5, bold=True, color=muted)
            set_run_font(finding_paragraph.add_run(_md(finding.get("finding"))), size=8.5, color=ink)
        if payload.get("message"):
            message_paragraph = doc.add_paragraph()
            set_run_font(message_paragraph.add_run("执行说明："), size=8.5, bold=True, color=muted)
            set_run_font(message_paragraph.add_run(_md(payload.get("message"))), size=8.5, color=ink)
        suggestions = finding.get("suggestions") or []
        if suggestions:
            suggestion_paragraph = doc.add_paragraph()
            set_run_font(suggestion_paragraph.add_run("处置建议："), size=8.5, bold=True, color=muted)
            set_run_font(
                suggestion_paragraph.add_run(
                    "；".join(str(value) for value in suggestions)
                    if isinstance(suggestions, list)
                    else str(suggestions)
                ),
                size=8.5,
                color=ink,
            )
        for sql in sql_list:
            sql_table = doc.add_table(rows=1, cols=1)
            sql_table.style = "Table Grid"
            sql_cell = sql_table.cell(0, 0)
            set_cell_text(sql_cell, sql, color="E7F8EF", size=7.5)
            for run in sql_cell.paragraphs[0].runs:
                set_run_font(run, size=7.5, color="E7F8EF", name="Consolas")
            shade_cell(sql_cell, "17202E")
            set_cell_margin(sql_cell, top=100, start=130, bottom=100, end=130)
            set_table_widths(sql_table, [25.2])
            prevent_row_split(sql_table.rows[0])
        if rows and columns:
            max_columns_per_table = 8
            for start in range(0, len(columns), max_columns_per_table):
                chunk = columns[start:start + max_columns_per_table]
                if len(columns) > max_columns_per_table:
                    field_label = doc.add_paragraph()
                    field_label.paragraph_format.space_before = Pt(4)
                    field_label.paragraph_format.space_after = Pt(3)
                    set_run_font(
                        field_label.add_run(
                            f"查询结果字段 {start + 1}-{start + len(chunk)} / {len(columns)}"
                        ),
                        size=7.5,
                        bold=True,
                        color=muted,
                    )
                chunk_rows = [[row.get(column) for column in chunk] for row in rows]
                widths = result_table_widths(chunk, chunk_rows)
                body_size = 7.5 if len(chunk) <= 5 else 6.8
                add_table(chunk, chunk_rows, widths, body_size=body_size)
        else:
            empty_paragraph = doc.add_paragraph()
            set_run_font(
                empty_paragraph.add_run(
                    f"查询结果：{_export_empty_result(status, payload)}"
                ),
                size=8,
                color=muted,
            )
        if payload.get("truncated"):
            truncated = doc.add_paragraph()
            set_run_font(
                truncated.add_run(
                    f"结果已截断，报告保留前 {payload.get('max_rows') or len(rows)} 行。"
                ),
                size=8,
                bold=True,
                color=orange,
            )
        add_spacer(3)

    ai_sections = (
        ("风险成因", ai_result.get("root_cause_analysis") or []),
        ("影响范围", ai_result.get("business_impact") or []),
        (
            "优先处置建议",
            [
                item.get("check_item") or item.get("finding")
                for item in ai_result.get("recommended_priorities") or []
            ],
        ),
    )
    has_ai_analysis = any(values for _, values in ai_sections)
    analysis_heading = add_section_heading("综合分析（规则引擎）")
    analysis_heading.paragraph_format.page_break_before = has_ai_analysis
    if has_ai_analysis:
        for label, values in ai_sections:
            add_section_heading(label, level=2)
            if values:
                for value in values:
                    paragraph = doc.add_paragraph(style="List Bullet")
                    set_run_font(paragraph.add_run(_md(value)), size=8.5, color=ink)
            else:
                paragraph = doc.add_paragraph()
                set_run_font(paragraph.add_run("暂无数据。"), size=8.5, color=muted)
    else:
        evidence_scope = "、".join(str(value) for value in ai_result.get("evidence_scope") or [])
        add_note(
            "分析结论",
            _export_no_risk_text(summary)
            + (f" 已纳入分析的证据范围：{evidence_scope}。" if evidence_scope else ""),
            fill="F1F8F4",
            accent=green,
        )

    add_section_heading("工单描述")
    add_note("建议内容", ai_result.get("ticket_description") or "-", fill=light_fill, accent=blue)

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream


@router.get("/{report_id}")
def get_report(report_id: int, db: Session = Depends(get_db)):
    report = db.query(DBCheckReport).filter(DBCheckReport.id == report_id).first()
    if not report:
        raise HTTPException(404, "报告不存在")
    result = json.loads(report.result_json) if report.result_json else {}
    result["report_id"] = report.id
    result["instance_name"] = report.instance_name
    return result


@router.delete("/{report_id}")
def delete_report(report_id: int, db: Session = Depends(get_db)):
    report = db.query(DBCheckReport).filter(DBCheckReport.id == report_id).first()
    if not report:
        raise HTTPException(404, "报告不存在")
    db.delete(report)
    db.commit()
    return {"detail": "报告已删除"}


@router.get("/{report_id}/export-word")
def export_word(report_id: int, db: Session = Depends(get_db)):
    report = db.query(DBCheckReport).filter(DBCheckReport.id == report_id).first()
    if not report:
        raise HTTPException(404, "报告不存在")

    data = json.loads(report.result_json) if report.result_json else {}
    items = data.get("items", [])

    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    title = doc.add_heading("", level=0)
    run = title.add_run("Oracle 数据库巡检报告")
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0, 51, 102)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    info_para = doc.add_paragraph()
    info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_run = info_para.add_run(
        f"实例：{report.instance_name}    |    "
        f"时间：{report.create_time.strftime('%Y-%m-%d %H:%M:%S') if report.create_time else '-'}    |    "
        f"检查项：{report.total}    通过：{report.success}    失败：{report.failed}"
    )
    info_run.font.size = Pt(10)
    info_run.font.color.rgb = RGBColor(102, 102, 102)

    doc.add_paragraph("")

    summary = doc.add_table(rows=4, cols=4, style="Light Grid Accent 1")
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    cells = summary.rows[0].cells
    cells[0].text = "检查项"
    cells[1].text = "通过"
    cells[2].text = "失败"
    cells[3].text = "通过率"
    cells = summary.rows[1].cells
    cells[0].text = str(report.total)
    cells[1].text = str(report.success)
    cells[2].text = str(report.failed)
    pct = f"{report.success / report.total * 100:.1f}%" if report.total > 0 else "0%"
    cells[3].text = pct
    for cell in summary.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True

    doc.add_paragraph("")

    for item in items:
        h = doc.add_heading("", level=2)
        status_text = "✓" if item.get("status") == "ok" else "✗"
        run = h.add_run(f"{status_text} {item.get('title', '')}")
        if item.get("status") == "ok":
            run.font.color.rgb = RGBColor(0, 128, 0)
        else:
            run.font.color.rgb = RGBColor(255, 0, 0)
        run.font.size = Pt(13)

        if item.get("status") == "error":
            p = doc.add_paragraph()
            r = p.add_run(f"错误：{item.get('error', '')}")
            r.font.color.rgb = RGBColor(255, 0, 0)
            continue

        columns = item.get("columns", [])
        rows = item.get("rows", [])
        if not columns or not rows:
            p = doc.add_paragraph("无数据")
            p.runs[0].font.color.rgb = RGBColor(153, 153, 153)
            continue

        max_rows = min(len(rows), 100)
        table = doc.add_table(rows=max_rows + 1, cols=len(columns), style="Light List Accent 1")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for j, col_name in enumerate(columns):
            cell = table.rows[0].cells[j]
            cell.text = col_name
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(8)

        for i, row in enumerate(rows[:max_rows]):
            for j, val in enumerate(row):
                cell = table.rows[i + 1].cells[j]
                cell.text = str(val) if val is not None else ""
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8)

        if len(rows) > 100:
            p = doc.add_paragraph(f"... 共 {len(rows)} 行，仅展示前 100 行")
            p.runs[0].font.color.rgb = RGBColor(153, 153, 153)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    filename = f"dbcheck_{report.instance_name}_{report.create_time.strftime('%Y%m%d_%H%M%S') if report.create_time else 'report'}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )
