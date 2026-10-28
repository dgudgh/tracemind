from __future__ import annotations

import io
import json
import re
import uuid
from datetime import datetime, time

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt
from fastapi import APIRouter, Depends, HTTPException, Query, Request
from pydantic import BaseModel, Field
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.models.db_instance import DBInstance
from app.models.report_center import (
    ReportArtifact,
    ReportArtifactAttachment,
    ReportSchedule,
    ReportScheduleRun,
    ReportSubscription,
)
from app.schemas.report_center import (
    ReportArchiveRequest,
    ReportAttachmentCreate,
    ReportLegalHoldRequest,
    ReportScheduleCreate,
    ReportScheduleUpdate,
    ReportSnapshotCreate,
    ReportSubscriptionCreate,
    ReportSubscriptionUpdate,
)
from app.services import report_center, report_document
from app.services.access_control import resolve_scope


router = APIRouter(prefix="/reports", tags=["统一报告中心"])


def _auth(request: Request) -> tuple[int, str, bool]:
    info = getattr(request.state, "auth", {}) or {}
    return int(info.get("user_id") or 0), str(info.get("username") or "system"), info.get("role") == "admin"


def _allowed_instance_ids(request: Request, db: Session, minimum_access: str = "read") -> set[int] | None:
    user_id, _, is_admin = _auth(request)
    if is_admin:
        return None
    scope = resolve_scope(db, user_id, minimum_access=minimum_access)
    return set(scope.native_instance_ids) if scope.restricted else None


def _require_instance_scope(request: Request, db: Session, instance_ids: set[int], minimum_access: str = "read") -> None:
    allowed = _allowed_instance_ids(request, db, minimum_access)
    if allowed is not None and not instance_ids.issubset(allowed):
        raise HTTPException(403, "报告包含当前用户数据范围之外的实例")


def _artifact_or_404(db: Session, artifact_id: int) -> ReportArtifact:
    row = db.query(ReportArtifact).filter(ReportArtifact.id == artifact_id).first()
    if row is None:
        raise HTTPException(404, "报告制品不存在")
    return row


def _authorize_artifact(request: Request, db: Session, row: ReportArtifact, minimum_access: str = "read") -> None:
    allowed = _allowed_instance_ids(request, db, minimum_access)
    scope = report_center.artifact_instance_scope(row)
    if allowed is not None and (not scope or not scope.issubset(allowed)):
        raise HTTPException(403, "报告不在当前用户的数据范围内")


def _schedule_or_404(db: Session, schedule_id: int) -> ReportSchedule:
    row = db.query(ReportSchedule).filter(ReportSchedule.id == schedule_id).first()
    if row is None:
        raise HTTPException(404, "报告调度不存在")
    return row


def _schedule_instance_ids(row_or_values) -> set[int]:
    filter_json = row_or_values.filter_json if hasattr(row_or_values, "filter_json") else row_or_values.get("filter_json") or {}
    return {int(value) for value in filter_json.get("instance_ids", []) if str(value).isdigit()}


def _authorize_schedule(request: Request, db: Session, row: ReportSchedule, minimum_access: str = "read") -> None:
    allowed = _allowed_instance_ids(request, db, minimum_access)
    if allowed is None:
        return
    scoped = _schedule_instance_ids(row)
    if not scoped or not scoped.issubset(allowed):
        raise HTTPException(403, "调度范围超出当前用户的数据权限")


def _parse_date(value: str | None, *, end_of_day: bool = False) -> datetime | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(value)
    except ValueError as exc:
        raise HTTPException(422, "日期格式必须为 ISO 8601") from exc
    if end_of_day and len(value) <= 10:
        parsed = datetime.combine(parsed.date(), time.max)
    # Date-only and naive catalog inputs refer to the displayed Shanghai wall time.
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=report_center.CATALOG_TIMEZONE)


def _collect_and_sync(
    db: Session,
    actor: str,
    report_type: str = "all",
    limit: int = 500,
    allowed_instance_ids: set[int] | None = None,
) -> list[ReportArtifact]:
    from app.api.v1.report import _collect_unified_reports

    sources = _collect_unified_reports(db, report_type, limit)
    if allowed_instance_ids is not None:
        sources = [source for source in sources if source.get("instance_id") in allowed_instance_ids]
    return report_center.sync_source_artifacts(db, sources, actor=actor)


@router.post("/catalog/sync")
def sync_report_catalog(request: Request, limit: int = Query(500, ge=1, le=1000), db: Session = Depends(get_db)):
    _, actor, _ = _auth(request)
    allowed = _allowed_instance_ids(request, db)
    rows = _collect_and_sync(db, actor, "all", limit, allowed)
    return {"synced": len(rows), "items": [report_center.serialize_artifact(row) for row in rows]}


@router.get("/catalog")
def list_report_catalog(
    request: Request,
    report_type: str = Query("all"),
    keyword: str | None = Query(None, max_length=200),
    instance_id: int | None = Query(None),
    archive_state: str = Query("all", pattern="^(all|active|archived)$"),
    date_from: str | None = Query(None),
    date_to: str | None = Query(None),
    refresh_sources: bool = Query(False),
    limit: int = Query(200, ge=1, le=1000),
    db: Session = Depends(get_db),
):
    _, actor, _ = _auth(request)
    allowed = _allowed_instance_ids(request, db)
    if refresh_sources:
        _collect_and_sync(db, actor, report_type, min(limit, 500), allowed)
    rows = report_center.list_artifacts(
        db,
        report_type=report_type,
        keyword=keyword,
        instance_id=instance_id,
        archive_state=archive_state,
        date_from=_parse_date(date_from),
        date_to=_parse_date(date_to, end_of_day=True),
        allowed_instance_ids=allowed,
        limit=limit,
        metadata_only=True,
    )
    return {"items": [report_center.serialize_artifact(row) for row in rows]}


@router.get("/catalog/stats")
def report_catalog_stats(request: Request, db: Session = Depends(get_db)):
    rows = report_center.list_artifacts(db, allowed_instance_ids=_allowed_instance_ids(request, db), limit=1000, metadata_only=True)
    return report_center.artifact_stats(rows)


class ReportDocumentGenerateRequest(BaseModel):
    instance_id: int = Field(gt=0)
    report_type: str = Field(default="both", pattern="^(unified|after_sales|both)$")
    window_hours: int = Field(default=168, ge=1, le=720)
    format: str = Field(default="markdown", pattern="^(markdown|html)$")
    include_ai: bool = True


@router.post("/document/generate")
def generate_report_document(payload: ReportDocumentGenerateRequest, request: Request, db: Session = Depends(get_db)):
    """One-click derived report generation for the selected instance (no target DB writes)."""
    _require_instance_scope(request, db, {payload.instance_id}, "read")
    instance = db.get(DBInstance, payload.instance_id)
    if instance is None:
        raise HTTPException(404, "实例不存在")
    _, actor, _ = _auth(request)
    types = ["unified", "after_sales"] if payload.report_type == "both" else [payload.report_type]
    items = []
    try:
        for report_type in types:
            item = report_document.generate(
                db, instance, report_type=report_type, window_hours=payload.window_hours,
                format=payload.format, include_ai=payload.include_ai,
                execution_id=f"ui:{uuid.uuid4().hex}", actor=actor, commit=False,
            )
            items.append(item)
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"schema": "report.document_batch/v1", "items": items}


@router.post("/catalog/snapshots")
def create_report_snapshot(payload: ReportSnapshotCreate, request: Request, db: Session = Depends(get_db)):
    _, actor, _ = _auth(request)
    requested_ids = set(payload.instance_ids)
    _require_instance_scope(request, db, requested_ids, "operate")
    from app.api.v1.report import _collect_unified_reports

    items = _collect_unified_reports(db, payload.report_type, payload.limit)
    allowed = _allowed_instance_ids(request, db, "operate")
    if requested_ids:
        items = [item for item in items if item.get("instance_id") in requested_ids]
    elif allowed is not None:
        items = [item for item in items if item.get("instance_id") in allowed]
    now = report_center.utcnow()
    artifact = report_center.create_snapshot_artifact(
        db,
        items,
        source_key=f"manual:{uuid.uuid4().hex}",
        title=payload.title,
        report_type=payload.report_type,
        actor=actor,
        watermark=payload.watermark,
        retention_days=payload.retention_days,
        generated_at=now,
    )
    return report_center.serialize_artifact(artifact)


@router.get("/catalog/{artifact_id}")
def report_artifact_detail(artifact_id: int, request: Request, db: Session = Depends(get_db)):
    row = _artifact_or_404(db, artifact_id)
    _authorize_artifact(request, db, row)
    attachments = db.query(ReportArtifactAttachment).filter(ReportArtifactAttachment.artifact_id == row.id).order_by(ReportArtifactAttachment.id).all()
    result = report_center.serialize_artifact(row, include_content=True)
    result["attachments"] = [report_center.serialize_attachment(item) for item in attachments]
    result["verification"] = report_center.verify_artifact(row)
    return result


@router.get("/catalog/{artifact_id}/verify")
def verify_report_artifact(artifact_id: int, request: Request, db: Session = Depends(get_db)):
    row = _artifact_or_404(db, artifact_id)
    _authorize_artifact(request, db, row)
    return report_center.verify_artifact(row)


@router.post("/catalog/{artifact_id}/archive")
def archive_report_artifact(artifact_id: int, payload: ReportArchiveRequest, request: Request, db: Session = Depends(get_db)):
    row = _artifact_or_404(db, artifact_id)
    _authorize_artifact(request, db, row, "operate")
    _, actor, _ = _auth(request)
    return report_center.serialize_artifact(report_center.archive_artifact(db, row, actor=actor, archive_uri=payload.archive_uri, reason=payload.reason))


@router.put("/catalog/{artifact_id}/legal-hold")
def legal_hold_report_artifact(artifact_id: int, payload: ReportLegalHoldRequest, request: Request, db: Session = Depends(get_db)):
    row = _artifact_or_404(db, artifact_id)
    _authorize_artifact(request, db, row, "operate")
    _, actor, _ = _auth(request)
    return report_center.serialize_artifact(report_center.set_legal_hold(db, row, enabled=payload.enabled, actor=actor, reason=payload.reason))


@router.post("/catalog/{artifact_id}/attachments")
def create_report_attachment(artifact_id: int, payload: ReportAttachmentCreate, request: Request, db: Session = Depends(get_db)):
    row = _artifact_or_404(db, artifact_id)
    _authorize_artifact(request, db, row, "operate")
    attachment = report_center.add_attachment(db, row, payload.model_dump(exclude_none=True))
    return report_center.serialize_attachment(attachment)


@router.get("/catalog/{artifact_id}/download")
def download_report_artifact(
    artifact_id: int,
    request: Request,
    format: str = Query("markdown", pattern="^(markdown|html|json|word)$"),
    db: Session = Depends(get_db),
):
    row = _artifact_or_404(db, artifact_id)
    _authorize_artifact(request, db, row)
    verification = report_center.verify_artifact(row)
    if not verification["valid"]:
        raise HTTPException(409, "报告完整性校验失败，已拒绝下载")
    if format not in (row.format_manifest or []):
        raise HTTPException(409, f"该报告不支持 {format} 格式")
    if format == "markdown":
        content = (row.content_markdown or "").encode("utf-8")
        media_type, suffix = "text/markdown; charset=utf-8", "md"
    elif format == "html":
        content = (row.content_html or "").encode("utf-8")
        media_type, suffix = "text/html; charset=utf-8", "html"
    elif format == "json":
        content = json.dumps(row.content_json or {}, ensure_ascii=False, indent=2, default=str).encode("utf-8")
        media_type, suffix = "application/json; charset=utf-8", "json"
    else:
        content = _docx_bytes(row)
        media_type, suffix = "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
    safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", row.artifact_uid)
    return StreamingResponse(
        io.BytesIO(content),
        media_type=media_type,
        headers={
            "Content-Disposition": f"attachment; filename={safe_name}.{suffix}",
            "X-Report-SHA256": row.content_sha256 or "",
            "X-Report-Signature-Key": row.signing_key_id or "",
        },
    )


def _docx_bytes(row: ReportArtifact) -> bytes:
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    document.add_heading(row.title, level=0)
    document.add_paragraph(f"报告编号：{row.artifact_uid}")
    document.add_paragraph(f"内容摘要：{row.content_sha256 or '-'}")
    document.add_paragraph(f"签名密钥：{row.signing_key_id or '-'}")
    payload = row.content_json if isinstance(row.content_json, dict) else {}
    items = payload.get("items") if isinstance(payload, dict) else None
    if isinstance(items, list):
        table = document.add_table(rows=1, cols=6)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells, ["类型", "标题", "实例", "数据库", "状态", "风险"]):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for item in items:
            cells = table.add_row().cells
            values = [item.get("report_type"), item.get("title"), item.get("instance_name"), item.get("db_type"), item.get("status"), item.get("risk_level")]
            for cell, value in zip(cells, values):
                cell.text = str(value or "-")
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    else:
        document.add_heading("报告摘要", level=1)
        document.add_paragraph(row.summary or "暂无摘要。")
    if row.watermark:
        document.add_paragraph(f"水印：{row.watermark}")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


@router.get("/schedules")
def list_report_schedules(request: Request, db: Session = Depends(get_db)):
    rows = db.query(ReportSchedule).order_by(ReportSchedule.id.desc()).all()
    allowed = _allowed_instance_ids(request, db)
    if allowed is not None:
        rows = [row for row in rows if _schedule_instance_ids(row) and _schedule_instance_ids(row).issubset(allowed)]
    return {"items": [report_center.serialize_schedule(row) for row in rows]}


@router.post("/schedules")
def create_report_schedule(payload: ReportScheduleCreate, request: Request, db: Session = Depends(get_db)):
    values = payload.model_dump()
    schedule_instance_ids = _schedule_instance_ids(values)
    allowed = _allowed_instance_ids(request, db, "operate")
    if allowed is not None and not schedule_instance_ids:
        raise HTTPException(403, "受限用户创建调度时必须明确选择实例范围")
    _require_instance_scope(request, db, schedule_instance_ids, "operate")
    _, actor, _ = _auth(request)
    try:
        row = report_center.create_schedule(db, values, actor=actor)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    return report_center.serialize_schedule(row)


@router.put("/schedules/{schedule_id}")
def update_report_schedule(schedule_id: int, payload: ReportScheduleUpdate, request: Request, db: Session = Depends(get_db)):
    row = _schedule_or_404(db, schedule_id)
    _authorize_schedule(request, db, row, "operate")
    values = payload.model_dump(exclude_unset=True)
    if "filter_json" in values:
        schedule_instance_ids = _schedule_instance_ids(values)
        allowed = _allowed_instance_ids(request, db, "operate")
        if allowed is not None and not schedule_instance_ids:
            raise HTTPException(403, "受限用户的调度必须保留明确的实例范围")
        _require_instance_scope(request, db, schedule_instance_ids, "operate")
    _, actor, _ = _auth(request)
    try:
        row = report_center.update_schedule(db, row, values, actor=actor)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    return report_center.serialize_schedule(row)


@router.delete("/schedules/{schedule_id}")
def delete_report_schedule(schedule_id: int, request: Request, db: Session = Depends(get_db)):
    row = _schedule_or_404(db, schedule_id)
    _authorize_schedule(request, db, row, "operate")
    _, actor, _ = _auth(request)
    from app.services.access_control import append_security_audit

    append_security_audit(db, category="report", action="schedule_delete", actor=actor, target_type="report_schedule", target_id=row.id, details={"name": row.name})
    db.delete(row)
    db.commit()
    return {"deleted": True}


@router.post("/schedules/{schedule_id}/run")
def run_report_schedule(schedule_id: int, request: Request, db: Session = Depends(get_db)):
    row = _schedule_or_404(db, schedule_id)
    _authorize_schedule(request, db, row, "operate")
    _, actor, _ = _auth(request)
    run = report_center.execute_schedule(db, row, actor=actor, trigger_type="manual")
    return report_center.serialize_run(run)


@router.get("/schedules/{schedule_id}/runs")
def list_report_schedule_runs(schedule_id: int, request: Request, limit: int = Query(100, ge=1, le=1000), db: Session = Depends(get_db)):
    row = _schedule_or_404(db, schedule_id)
    _authorize_schedule(request, db, row)
    runs = db.query(ReportScheduleRun).filter(ReportScheduleRun.schedule_id == schedule_id).order_by(ReportScheduleRun.planned_at.desc()).limit(limit).all()
    return {"items": [report_center.serialize_run(item) for item in runs]}


@router.get("/subscriptions")
def list_report_subscriptions(request: Request, db: Session = Depends(get_db)):
    rows = db.query(ReportSubscription).order_by(ReportSubscription.id.desc()).all()
    allowed_schedule_ids = {item["id"] for item in list_report_schedules(request, db)["items"]}
    allowed = _allowed_instance_ids(request, db)
    if allowed is not None:
        rows = [row for row in rows if row.schedule_id in allowed_schedule_ids]
    return {"items": [report_center.serialize_subscription(row) for row in rows]}


@router.post("/subscriptions")
def create_report_subscription(payload: ReportSubscriptionCreate, request: Request, db: Session = Depends(get_db)):
    if _allowed_instance_ids(request, db, "operate") is not None and payload.schedule_id is None:
        raise HTTPException(403, "受限用户必须将报告订阅关联到有权访问的调度策略")
    if payload.schedule_id is not None:
        _authorize_schedule(request, db, _schedule_or_404(db, payload.schedule_id), "operate")
    _, actor, _ = _auth(request)
    return report_center.serialize_subscription(report_center.create_subscription(db, payload.model_dump(), actor=actor))


@router.put("/subscriptions/{subscription_id}")
def update_report_subscription(subscription_id: int, payload: ReportSubscriptionUpdate, request: Request, db: Session = Depends(get_db)):
    row = db.query(ReportSubscription).filter(ReportSubscription.id == subscription_id).first()
    if row is None:
        raise HTTPException(404, "报告订阅不存在")
    if row.schedule_id is None and _allowed_instance_ids(request, db, "operate") is not None:
        raise HTTPException(403, "全局报告订阅仅平台管理员可管理")
    if row.schedule_id is not None:
        _authorize_schedule(request, db, _schedule_or_404(db, row.schedule_id), "operate")
    values = payload.model_dump(exclude_unset=True)
    if values.get("schedule_id") is not None:
        _authorize_schedule(request, db, _schedule_or_404(db, values["schedule_id"]), "operate")
    _, actor, _ = _auth(request)
    return report_center.serialize_subscription(report_center.update_subscription(db, row, values, actor=actor))


@router.delete("/subscriptions/{subscription_id}")
def delete_report_subscription(subscription_id: int, request: Request, db: Session = Depends(get_db)):
    row = db.query(ReportSubscription).filter(ReportSubscription.id == subscription_id).first()
    if row is None:
        raise HTTPException(404, "报告订阅不存在")
    if row.schedule_id is None and _allowed_instance_ids(request, db, "operate") is not None:
        raise HTTPException(403, "全局报告订阅仅平台管理员可管理")
    if row.schedule_id is not None:
        _authorize_schedule(request, db, _schedule_or_404(db, row.schedule_id), "operate")
    db.delete(row)
    db.commit()
    return {"deleted": True}
