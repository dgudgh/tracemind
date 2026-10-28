#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AWR报告导出模块
将分析结果和建议导出为多种格式

作者: 青学会MOP技术社区
日期: 2025-04-16
"""

import os
import logging
import markdown
import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

# 设置日志
logger = logging.getLogger("青学会MOP技术社区")
logger.setLevel(logging.INFO)
if not logger.handlers:
    handler = logging.StreamHandler()
    formatter = logging.Formatter('%(asctime)s - %(name)s - %(levelname)s - %(message)s')
    handler.setFormatter(formatter)
    logger.addHandler(handler)

class AWRReportExporter:
    """
    AWR报告导出类
    将分析结果和建议导出为多种格式
    """
    
    def __init__(self, analysis_results, recommendations):
        """
        初始化导出器
        
        参数:
            analysis_results (dict): 分析结果
            recommendations (dict): 建议结果
        """
        self.analysis_results = analysis_results
        self.recommendations = recommendations
        self.timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        logger.info("初始化AWR报告导出器")
    
    def export(self, format_type, output_dir, filename=None):
        """
        导出报告
        
        参数:
            format_type (str): 导出格式，支持 'md', 'pdf', 'txt', 'doc'
            output_dir (str): 输出目录
            filename (str, optional): 输出文件名，不包含扩展名
            
        返回:
            str: 导出文件的完整路径
        """
        # 确保输出目录存在
        os.makedirs(output_dir, exist_ok=True)
        
        # 如果未指定文件名，使用默认文件名
        if not filename:
            db_name = self.analysis_results.get('summary', {}).get('database_name', 'Unknown')
            filename = f"AWR_Analysis_{db_name}_{self.timestamp}"
        
        # 根据格式类型调用相应的导出方法
        if format_type.lower() == 'md':
            return self._export_md(output_dir, filename)
        elif format_type.lower() == 'pdf':
            return self._export_pdf(output_dir, filename)
        elif format_type.lower() == 'txt':
            return self._export_txt(output_dir, filename)
        elif format_type.lower() == 'doc':
            return self._export_doc(output_dir, filename)
        else:
            logger.error(f"不支持的导出格式: {format_type}")
            return None
    
    def _export_md(self, output_dir, filename):
        """导出为Markdown格式"""
        try:
            # 构建输出文件路径
            output_file = os.path.join(output_dir, f"{filename}.md")
            
            # 生成Markdown内容
            md_content = self._generate_md_content()
            
            # 写入文件
            with open(output_file, 'w', encoding='utf-8') as f:
                f.write(md_content)
            
            logger.info(f"成功导出Markdown报告: {output_file}")
            return output_file
        except Exception as e:
            logger.error(f"导出Markdown报告失败: {str(e)}")
            return None
    
    def _export_pdf(self, output_dir, filename):
        """导出为PDF格式"""
        try:
            # 构建输出文件路径
            output_file = os.path.join(output_dir, f"{filename}.pdf")
            
            # 创建PDF文档
            doc = SimpleDocTemplate(output_file, pagesize=A4)
            
            # 生成PDF内容
            pdf_content = self._generate_pdf_content()
            
            # 构建PDF
            doc.build(pdf_content)
            
            logger.info(f"成功导出PDF报告: {output_file}")
            return output_file
        except Exception as e:
            logger.error(f"导出PDF报告失败: {str(e)}")
            return None
    
    def _export_txt(self, output_dir, filename):
        """导出为纯文本格式"""
        try:
            # 构建输出文件路径
            output_file = os.path.join(output_dir, f"{filename}.txt")
            
            # 生成文本内容
            txt_content = self._generate_txt_content()
            
            # 写入文件
            with open(output_file, 'w', encoding='utf-8') as f:
                f.write(txt_content)
            
            logger.info(f"成功导出文本报告: {output_file}")
            return output_file
        except Exception as e:
            logger.error(f"导出文本报告失败: {str(e)}")
            return None
    
    def _export_doc(self, output_dir, filename):
        """导出为Word文档格式"""
        try:
            # 构建输出文件路径
            output_file = os.path.join(output_dir, f"{filename}.docx")
            
            # 创建Word文档
            doc = Document()
            
            # 生成Word文档内容
            self._generate_doc_content(doc)
            
            # 保存文档
            doc.save(output_file)
            
            logger.info(f"成功导出Word报告: {output_file}")
            return output_file
        except Exception as e:
            logger.error(f"导出Word报告失败: {str(e)}")
            return None
    
    def _generate_md_content(self):
        """生成Markdown内容"""
        md = []
        
        # 添加标题
        summary = self.recommendations.get('summary', {})
        db_name = summary.get('database_name', 'Unknown')
        instance_name = summary.get('instance_name', 'Unknown')
        
        md.append(f"# Oracle AWR报告分析 - {db_name}/{instance_name}\n")
        md.append(f"*生成时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}*\n")
        md.append("*由青学会MOP技术社区AWR分析工具生成*\n\n")
        
        # 添加概要信息
        md.append("## 1. 分析概要\n")
        md.append(f"- **数据库名称**: {summary.get('database_name', 'Unknown')}")
        md.append(f"- **实例名称**: {summary.get('instance_name', 'Unknown')}")
        md.append(f"- **数据库版本**: {summary.get('version', 'Unknown')}")
        md.append(f"- **主机名**: {summary.get('host', 'Unknown')}")
        md.append(f"- **快照时间段**: {summary.get('snapshot_period', 'Unknown')}")
        md.append(f"- **性能评分**: {summary.get('performance_score', 0)}/100 ({summary.get('performance_level', 'Unknown')})")
        md.append(f"- **整体评估**: {summary.get('overall_assessment', '')}\n")
        
        # 添加优先建议
        md.append("## 2. 优先建议\n")
        priority_recommendations = self.recommendations.get('priority_recommendations', [])
        
        if priority_recommendations:
            for i, rec in enumerate(priority_recommendations, 1):
                md.append(f"### 2.{i} {rec.get('title', 'Unknown')} ({rec.get('severity', 'Unknown')})\n")
                md.append(f"**发现**: {rec.get('finding', '')}\n")
                md.append(f"**影响**: {rec.get('impact', '')}\n")
                md.append("**建议**:")
                for r in rec.get('recommendations', []):
                    md.append(f"- {r}")
                md.append("")
        else:
            md.append("*没有优先建议*\n")
        
        # 添加CPU分析
        md.append("## 3. CPU分析\n")
        cpu_analysis = self.analysis_results.get('cpu_analysis', {})
        
        md.append(f"- **CPU核心数**: {cpu_analysis.get('cpu_cores', 0)}")
        
        host_cpu = cpu_analysis.get('host_cpu_usage', {})
        md.append(f"- **主机CPU使用率**:")
        md.append(f"  - 用户: {host_cpu.get('user', 0)}%")
        md.append(f"  - 系统: {host_cpu.get('system', 0)}%")
        md.append(f"  - 空闲: {host_cpu.get('idle', 0)}%")
        md.append(f"  - 总计: {host_cpu.get('total', 0)}%")
        
        db_cpu = cpu_analysis.get('db_cpu_usage', {})
        md.append(f"- **数据库CPU使用率**:")
        md.append(f"  - 总CPU占用: {db_cpu.get('total_cpu_pct', 0)}%")
        md.append(f"  - 每秒CPU时间: {db_cpu.get('per_second', 0)}秒")
        
        md.append(f"- **CPU利用率**: {cpu_analysis.get('cpu_utilization', 0)}%")
        md.append(f"- **CPU瓶颈**: {'是' if cpu_analysis.get('cpu_bottleneck', False) else '否'}\n")
        
        # 添加内存分析
        md.append("## 4. 内存分析\n")
        memory_analysis = self.analysis_results.get('memory_analysis', {})
        
        md.append(f"- **缓冲区命中率**: {memory_analysis.get('buffer_hit_ratio', 0)}%")
        md.append(f"- **库缓存命中率**: {memory_analysis.get('library_hit_ratio', 0)}%")
        md.append(f"- **内存排序百分比**: {memory_analysis.get('in_memory_sort_pct', 0)}%")
        
        memory_issues = memory_analysis.get('memory_issues', {})
        md.append("- **内存问题**:")
        md.append(f"  - 低缓冲区命中: {'是' if memory_issues.get('low_buffer_hit', False) else '否'}")
        md.append(f"  - 低库缓存命中: {'是' if memory_issues.get('low_library_hit', False) else '否'}")
        md.append(f"  - 低内存排序: {'是' if memory_issues.get('low_memory_sort', False) else '否'}\n")
        
        # 添加IO分析
        md.append("## 5. IO分析\n")
        io_analysis = self.analysis_results.get('io_analysis', {})
        
        md.append(f"- **每秒物理读**: {io_analysis.get('physical_reads_per_sec', 0)}")
        md.append(f"- **每秒物理写**: {io_analysis.get('physical_writes_per_sec', 0)}")
        md.append(f"- **每秒读IO请求**: {io_analysis.get('read_io_requests_per_sec', 0)}")
        md.append(f"- **每秒写IO请求**: {io_analysis.get('write_io_requests_per_sec', 0)}")
        md.append(f"- **每秒总IO请求**: {io_analysis.get('total_io_requests_per_sec', 0)}")
        md.append(f"- **读写比例**: {io_analysis.get('read_write_ratio', 0)}:1")
        md.append(f"- **IO瓶颈**: {'是' if io_analysis.get('io_bottleneck', False) else '否'}\n")
        
        # 添加等待事件分析
        md.append("## 6. 等待事件分析\n")
        wait_analysis = self.analysis_results.get('wait_analysis', {})
        
        md.append("### 6.1 前5个耗时事件\n")
        top_events = wait_analysis.get('top_events', [])
        
        if top_events:
            md.append("| 事件名称 | 等待时间(秒) | 数据库时间百分比 | 等待类别 |")
            md.append("| --- | --- | --- | --- |")
            for event in top_events:
                md.append(f"| {event.get('event_name', '')} | {event.get('wait_time', 0)} | {event.get('db_time_pct', 0)}% | {event.get('wait_class', '')} |")
        else:
            md.append("*没有等待事件数据*")
        
        md.append("\n### 6.2 等待类别分布\n")
        wait_class_distribution = wait_analysis.get('wait_class_distribution', {})
        
        if wait_class_distribution:
            md.append("| 等待类别 | 数据库时间百分比 |")
            md.append("| --- | --- |")
            for wait_class, pct in wait_class_distribution.items():
                md.append(f"| {wait_class} | {pct}% |")
        else:
            md.append("*没有等待类别分布数据*")
        
        md.append("\n### 6.3 等待事件特征\n")
        md.append(f"- **CPU绑定**: {'是' if wait_analysis.get('cpu_bound', False) else '否'}")
        md.append(f"- **IO绑定**: {'是' if wait_analysis.get('io_bound', False) else '否'}")
        md.append(f"- **并发问题**: {'是' if wait_analysis.get('concurrency_issues', False) else '否'}")
        md.append(f"- **网络问题**: {'是' if wait_analysis.get('network_issues', False) else '否'}\n")
        
        # 添加SQL分析
        md.append("## 7. SQL分析\n")
        sql_analysis = self.analysis_results.get('sql_analysis', {})
        
        md.append("### 7.1 高CPU消耗SQL\n")
        top_cpu_sql = sql_analysis.get('top_cpu_sql', [])
        
        if top_cpu_sql:
            for i, sql in enumerate(top_cpu_sql[:3], 1):
                md.append(f"**SQL {i}** (ID: {sql.get('sql_id', '')})")
                md.append(f"- CPU时间: {sql.get('cpu_time', 0)}秒")
                md.append(f"- 执行次数: {sql.get('executions', 0)}")
                md.append(f"- 每次执行CPU: {sql.get('cpu_per_exec', 0)}秒")
                md.append(f"- SQL文本: ```{sql.get('sql_text', '')[:200]}...```\n")
        else:
            md.append("*没有SQL分析数据*\n")
        
        # 添加实例效率分析
        md.append("## 8. 实例效率分析\n")
        instance_efficiency_analysis = self.analysis_results.get('instance_efficiency_analysis', {})
        
        md.append(f"- **缓冲区命中率**: {instance_efficiency_analysis.get('buffer_hit', 0)}%")
        md.append(f"- **缓冲区无等待率**: {instance_efficiency_analysis.get('buffer_nowait', 0)}%")
        md.append(f"- **库缓存命中率**: {instance_efficiency_analysis.get('library_hit', 0)}%")
        md.append(f"- **执行解析比率**: {instance_efficiency_analysis.get('execute_to_parse', 0)}%")
        md.append(f"- **软解析率**: {instance_efficiency_analysis.get('soft_parse', 0)}%")
        md.append(f"- **Latch命中率**: {instance_efficiency_analysis.get('latch_hit', 0)}%\n")
        
        # 添加潜在问题
        md.append("## 9. 潜在问题\n")
        potential_issues = self.analysis_results.get('potential_issues', [])
        
        if potential_issues:
            for i, issue in enumerate(potential_issues, 1):
                md.append(f"### 9.{i} {issue.get('issue', 'Unknown')} ({issue.get('severity', 'Unknown')})\n")
                md.append(f"**类别**: {issue.get('category', '')}")
                md.append(f"**描述**: {issue.get('description', '')}")
                md.append(f"**建议**: {issue.get('recommendation', '')}\n")
        else:
            md.append("*没有发现潜在问题*\n")
        
        # 添加一般性建议
        md.append("## 10. 一般性建议\n")
        general_recommendations = self.recommendations.get('general_recommendations', [])
        
        if general_recommendations:
            for i, rec in enumerate(general_recommendations, 1):
                md.append(f"### 10.{i} {rec.get('title', 'Unknown')}\n")
                md.append(f"**发现**: {rec.get('finding', '')}\n")
                md.append(f"**影响**: {rec.get('impact', '')}\n")
                md.append("**建议**:")
                for r in rec.get('recommendations', []):
                    md.append(f"- {r}")
                md.append("")
        else:
            md.append("*没有一般性建议*\n")
        
        # 添加结论
        md.append("## 11. 结论\n")
        performance_score = summary.get('performance_score', 0)
        performance_level = summary.get('performance_level', 'Unknown')
        
        md.append(f"根据AWR报告分析，数据库性能评分为**{performance_score}/100**，性能等级为**{performance_level}**。")
        md.append(f"{summary.get('overall_assessment', '')}")
        
        if performance_score < 60:
            md.append("建议优先解决上述优先建议中的高严重性问题，以提高数据库性能。")
        elif performance_score < 80:
            md.append("建议按照优先级顺序解决上述问题，以进一步提高数据库性能。")
        else:
            md.append("数据库性能良好，建议定期监控并进行预防性维护。")
        
        # 返回Markdown内容
        return "\n".join(md)
    
    def _generate_pdf_content(self):
        """生成PDF内容"""
        # 获取样式
        styles = getSampleStyleSheet()
        
        # 创建自定义样式
        title_style = ParagraphStyle(
            'Title',
            parent=styles['Title'],
            fontSize=16,
            spaceAfter=12
        )
        
        heading1_style = ParagraphStyle(
            'Heading1',
            parent=styles['Heading1'],
            fontSize=14,
            spaceAfter=10,
            spaceBefore=10
        )
        
        heading2_style = ParagraphStyle(
            'Heading2',
            parent=styles['Heading2'],
            fontSize=12,
            spaceAfter=8,
            spaceBefore=8
        )
        
        normal_style = styles['Normal']
        
        # 创建内容列表
        content = []
        
        # 添加标题
        summary = self.recommendations.get('summary', {})
        db_name = summary.get('database_name', 'Unknown')
        instance_name = summary.get('instance_name', 'Unknown')
        
        content.append(Paragraph(f"Oracle AWR报告分析 - {db_name}/{instance_name}", title_style))
        content.append(Paragraph(f"生成时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", normal_style))
        content.append(Paragraph("由青学会MOP技术社区AWR分析工具生成", normal_style))
        content.append(Spacer(1, 0.2*inch))
        
        # 添加概要信息
        content.append(Paragraph("1. 分析概要", heading1_style))
        
        summary_data = [
            ["数据库名称", summary.get('database_name', 'Unknown')],
            ["实例名称", summary.get('instance_name', 'Unknown')],
            ["数据库版本", summary.get('version', 'Unknown')],
            ["主机名", summary.get('host', 'Unknown')],
            ["快照时间段", summary.get('snapshot_period', 'Unknown')],
            ["性能评分", f"{summary.get('performance_score', 0)}/100 ({summary.get('performance_level', 'Unknown')})"],
            ["整体评估", summary.get('overall_assessment', '')]
        ]
        
        summary_table = Table(summary_data, colWidths=[2*inch, 4*inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('PADDING', (0, 0), (-1, -1), 6)
        ]))
        
        content.append(summary_table)
        content.append(Spacer(1, 0.2*inch))
        
        # 添加优先建议
        content.append(Paragraph("2. 优先建议", heading1_style))
        priority_recommendations = self.recommendations.get('priority_recommendations', [])
        
        if priority_recommendations:
            for i, rec in enumerate(priority_recommendations, 1):
                content.append(Paragraph(f"2.{i} {rec.get('title', 'Unknown')} ({rec.get('severity', 'Unknown')})", heading2_style))
                content.append(Paragraph(f"发现: {rec.get('finding', '')}", normal_style))
                content.append(Paragraph(f"影响: {rec.get('impact', '')}", normal_style))
                content.append(Paragraph("建议:", normal_style))
                
                for r in rec.get('recommendations', []):
                    content.append(Paragraph(f"• {r}", normal_style))
                
                content.append(Spacer(1, 0.1*inch))
        else:
            content.append(Paragraph("没有优先建议", normal_style))
        
        content.append(Spacer(1, 0.2*inch))
        
        # 添加CPU分析
        content.append(Paragraph("3. CPU分析", heading1_style))
        cpu_analysis = self.analysis_results.get('cpu_analysis', {})
        
        cpu_data = [
            ["CPU核心数", str(cpu_analysis.get('cpu_cores', 0))],
            ["主机用户CPU", f"{cpu_analysis.get('host_cpu_usage', {}).get('user', 0)}%"],
            ["主机系统CPU", f"{cpu_analysis.get('host_cpu_usage', {}).get('system', 0)}%"],
            ["主机空闲CPU", f"{cpu_analysis.get('host_cpu_usage', {}).get('idle', 0)}%"],
            ["主机总CPU使用率", f"{cpu_analysis.get('host_cpu_usage', {}).get('total', 0)}%"],
            ["数据库CPU占用", f"{cpu_analysis.get('db_cpu_usage', {}).get('total_cpu_pct', 0)}%"],
            ["每秒CPU时间", f"{cpu_analysis.get('db_cpu_usage', {}).get('per_second', 0)}秒"],
            ["CPU利用率", f"{cpu_analysis.get('cpu_utilization', 0)}%"],
            ["CPU瓶颈", "是" if cpu_analysis.get('cpu_bottleneck', False) else "否"]
        ]
        
        cpu_table = Table(cpu_data, colWidths=[2*inch, 4*inch])
        cpu_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('PADDING', (0, 0), (-1, -1), 6)
        ]))
        
        content.append(cpu_table)
        content.append(Spacer(1, 0.2*inch))
        
        # 添加内存分析
        content.append(Paragraph("4. 内存分析", heading1_style))
        memory_analysis = self.analysis_results.get('memory_analysis', {})
        
        memory_data = [
            ["缓冲区命中率", f"{memory_analysis.get('buffer_hit_ratio', 0)}%"],
            ["库缓存命中率", f"{memory_analysis.get('library_hit_ratio', 0)}%"],
            ["内存排序百分比", f"{memory_analysis.get('in_memory_sort_pct', 0)}%"],
            ["低缓冲区命中", "是" if memory_analysis.get('memory_issues', {}).get('low_buffer_hit', False) else "否"],
            ["低库缓存命中", "是" if memory_analysis.get('memory_issues', {}).get('low_library_hit', False) else "否"],
            ["低内存排序", "是" if memory_analysis.get('memory_issues', {}).get('low_memory_sort', False) else "否"]
        ]
        
        memory_table = Table(memory_data, colWidths=[2*inch, 4*inch])
        memory_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('PADDING', (0, 0), (-1, -1), 6)
        ]))
        
        content.append(memory_table)
        content.append(Spacer(1, 0.2*inch))
        
        # 添加IO分析
        content.append(Paragraph("5. IO分析", heading1_style))
        io_analysis = self.analysis_results.get('io_analysis', {})
        
        io_data = [
            ["每秒物理读", str(io_analysis.get('physical_reads_per_sec', 0))],
            ["每秒物理写", str(io_analysis.get('physical_writes_per_sec', 0))],
            ["每秒读IO请求", str(io_analysis.get('read_io_requests_per_sec', 0))],
            ["每秒写IO请求", str(io_analysis.get('write_io_requests_per_sec', 0))],
            ["每秒总IO请求", str(io_analysis.get('total_io_requests_per_sec', 0))],
            ["读写比例", f"{io_analysis.get('read_write_ratio', 0)}:1"],
            ["IO瓶颈", "是" if io_analysis.get('io_bottleneck', False) else "否"]
        ]
        
        io_table = Table(io_data, colWidths=[2*inch, 4*inch])
        io_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('PADDING', (0, 0), (-1, -1), 6)
        ]))
        
        content.append(io_table)
        content.append(Spacer(1, 0.2*inch))
        
        # 添加等待事件分析
        content.append(Paragraph("6. 等待事件分析", heading1_style))
        wait_analysis = self.analysis_results.get('wait_analysis', {})
        
        content.append(Paragraph("6.1 前5个耗时事件", heading2_style))
        top_events = wait_analysis.get('top_events', [])
        
        if top_events:
            event_data = [["事件名称", "等待时间(秒)", "数据库时间百分比", "等待类别"]]
            
            for event in top_events:
                event_data.append([
                    event.get('event_name', ''),
                    str(event.get('wait_time', 0)),
                    f"{event.get('db_time_pct', 0)}%",
                    event.get('wait_class', '')
                ])
            
            event_table = Table(event_data, colWidths=[2*inch, 1*inch, 1.5*inch, 1.5*inch])
            event_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('PADDING', (0, 0), (-1, -1), 6)
            ]))
            
            content.append(event_table)
        else:
            content.append(Paragraph("没有等待事件数据", normal_style))
        
        content.append(Spacer(1, 0.1*inch))
        
        content.append(Paragraph("6.2 等待类别分布", heading2_style))
        wait_class_distribution = wait_analysis.get('wait_class_distribution', {})
        
        if wait_class_distribution:
            wait_class_data = [["等待类别", "数据库时间百分比"]]
            
            for wait_class, pct in wait_class_distribution.items():
                wait_class_data.append([wait_class, f"{pct}%"])
            
            wait_class_table = Table(wait_class_data, colWidths=[3*inch, 3*inch])
            wait_class_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('PADDING', (0, 0), (-1, -1), 6)
            ]))
            
            content.append(wait_class_table)
        else:
            content.append(Paragraph("没有等待类别分布数据", normal_style))
        
        content.append(Spacer(1, 0.1*inch))
        
        content.append(Paragraph("6.3 等待事件特征", heading2_style))
        
        wait_feature_data = [
            ["CPU绑定", "是" if wait_analysis.get('cpu_bound', False) else "否"],
            ["IO绑定", "是" if wait_analysis.get('io_bound', False) else "否"],
            ["并发问题", "是" if wait_analysis.get('concurrency_issues', False) else "否"],
            ["网络问题", "是" if wait_analysis.get('network_issues', False) else "否"]
        ]
        
        wait_feature_table = Table(wait_feature_data, colWidths=[2*inch, 4*inch])
        wait_feature_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('PADDING', (0, 0), (-1, -1), 6)
        ]))
        
        content.append(wait_feature_table)
        content.append(Spacer(1, 0.2*inch))
        
        # 添加结论
        content.append(Paragraph("7. 结论", heading1_style))
        performance_score = summary.get('performance_score', 0)
        performance_level = summary.get('performance_level', 'Unknown')
        
        conclusion_text = f"根据AWR报告分析，数据库性能评分为{performance_score}/100，性能等级为{performance_level}。"
        conclusion_text += f" {summary.get('overall_assessment', '')}"
        
        if performance_score < 60:
            conclusion_text += " 建议优先解决上述优先建议中的高严重性问题，以提高数据库性能。"
        elif performance_score < 80:
            conclusion_text += " 建议按照优先级顺序解决上述问题，以进一步提高数据库性能。"
        else:
            conclusion_text += " 数据库性能良好，建议定期监控并进行预防性维护。"
        
        content.append(Paragraph(conclusion_text, normal_style))
        
        # 返回PDF内容
        return content
    
    def _generate_txt_content(self):
        """生成纯文本内容"""
        txt = []
        
        # 添加标题
        summary = self.recommendations.get('summary', {})
        db_name = summary.get('database_name', 'Unknown')
        instance_name = summary.get('instance_name', 'Unknown')
        
        txt.append("=" * 80)
        txt.append(f"Oracle AWR报告分析 - {db_name}/{instance_name}")
        txt.append(f"生成时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        txt.append("由青学会MOP技术社区AWR分析工具生成")
        txt.append("=" * 80)
        txt.append("")
        
        # 添加概要信息
        txt.append("1. 分析概要")
        txt.append("-" * 80)
        txt.append(f"数据库名称: {summary.get('database_name', 'Unknown')}")
        txt.append(f"实例名称: {summary.get('instance_name', 'Unknown')}")
        txt.append(f"数据库版本: {summary.get('version', 'Unknown')}")
        txt.append(f"主机名: {summary.get('host', 'Unknown')}")
        txt.append(f"快照时间段: {summary.get('snapshot_period', 'Unknown')}")
        txt.append(f"性能评分: {summary.get('performance_score', 0)}/100 ({summary.get('performance_level', 'Unknown')})")
        txt.append(f"整体评估: {summary.get('overall_assessment', '')}")
        txt.append("")
        
        # 添加优先建议
        txt.append("2. 优先建议")
        txt.append("-" * 80)
        priority_recommendations = self.recommendations.get('priority_recommendations', [])
        
        if priority_recommendations:
            for i, rec in enumerate(priority_recommendations, 1):
                txt.append(f"2.{i} {rec.get('title', 'Unknown')} ({rec.get('severity', 'Unknown')})")
                txt.append("")
                txt.append(f"发现: {rec.get('finding', '')}")
                txt.append(f"影响: {rec.get('impact', '')}")
                txt.append("建议:")
                for r in rec.get('recommendations', []):
                    txt.append(f"- {r}")
                txt.append("")
        else:
            txt.append("没有优先建议")
            txt.append("")
        
        # 添加CPU分析
        txt.append("3. CPU分析")
        txt.append("-" * 80)
        cpu_analysis = self.analysis_results.get('cpu_analysis', {})
        
        txt.append(f"CPU核心数: {cpu_analysis.get('cpu_cores', 0)}")
        
        host_cpu = cpu_analysis.get('host_cpu_usage', {})
        txt.append("主机CPU使用率:")
        txt.append(f"  用户: {host_cpu.get('user', 0)}%")
        txt.append(f"  系统: {host_cpu.get('system', 0)}%")
        txt.append(f"  空闲: {host_cpu.get('idle', 0)}%")
        txt.append(f"  总计: {host_cpu.get('total', 0)}%")
        
        db_cpu = cpu_analysis.get('db_cpu_usage', {})
        txt.append("数据库CPU使用率:")
        txt.append(f"  总CPU占用: {db_cpu.get('total_cpu_pct', 0)}%")
        txt.append(f"  每秒CPU时间: {db_cpu.get('per_second', 0)}秒")
        
        txt.append(f"CPU利用率: {cpu_analysis.get('cpu_utilization', 0)}%")
        txt.append(f"CPU瓶颈: {'是' if cpu_analysis.get('cpu_bottleneck', False) else '否'}")
        txt.append("")
        
        # 添加内存分析
        txt.append("4. 内存分析")
        txt.append("-" * 80)
        memory_analysis = self.analysis_results.get('memory_analysis', {})
        
        txt.append(f"缓冲区命中率: {memory_analysis.get('buffer_hit_ratio', 0)}%")
        txt.append(f"库缓存命中率: {memory_analysis.get('library_hit_ratio', 0)}%")
        txt.append(f"内存排序百分比: {memory_analysis.get('in_memory_sort_pct', 0)}%")
        
        memory_issues = memory_analysis.get('memory_issues', {})
        txt.append("内存问题:")
        txt.append(f"  低缓冲区命中: {'是' if memory_issues.get('low_buffer_hit', False) else '否'}")
        txt.append(f"  低库缓存命中: {'是' if memory_issues.get('low_library_hit', False) else '否'}")
        txt.append(f"  低内存排序: {'是' if memory_issues.get('low_memory_sort', False) else '否'}")
        txt.append("")
        
        # 添加IO分析
        txt.append("5. IO分析")
        txt.append("-" * 80)
        io_analysis = self.analysis_results.get('io_analysis', {})
        
        txt.append(f"每秒物理读: {io_analysis.get('physical_reads_per_sec', 0)}")
        txt.append(f"每秒物理写: {io_analysis.get('physical_writes_per_sec', 0)}")
        txt.append(f"每秒读IO请求: {io_analysis.get('read_io_requests_per_sec', 0)}")
        txt.append(f"每秒写IO请求: {io_analysis.get('write_io_requests_per_sec', 0)}")
        txt.append(f"每秒总IO请求: {io_analysis.get('total_io_requests_per_sec', 0)}")
        txt.append(f"读写比例: {io_analysis.get('read_write_ratio', 0)}:1")
        txt.append(f"IO瓶颈: {'是' if io_analysis.get('io_bottleneck', False) else '否'}")
        txt.append("")
        
        # 添加等待事件分析
        txt.append("6. 等待事件分析")
        txt.append("-" * 80)
        wait_analysis = self.analysis_results.get('wait_analysis', {})
        
        txt.append("6.1 前5个耗时事件")
        txt.append("")
        top_events = wait_analysis.get('top_events', [])
        
        if top_events:
            # 创建表头
            event_header = f"{'事件名称':<30} {'等待时间(秒)':<15} {'数据库时间百分比':<20} {'等待类别':<15}"
            txt.append(event_header)
            txt.append("-" * 80)
            
            for event in top_events:
                event_row = f"{event.get('event_name', ''):<30} {str(event.get('wait_time', 0)):<15} {str(event.get('db_time_pct', 0))+'%':<20} {event.get('wait_class', ''):<15}"
                txt.append(event_row)
        else:
            txt.append("没有等待事件数据")
        
        txt.append("")
        txt.append("6.2 等待类别分布")
        txt.append("")
        wait_class_distribution = wait_analysis.get('wait_class_distribution', {})
        
        if wait_class_distribution:
            # 创建表头
            wait_class_header = f"{'等待类别':<30} {'数据库时间百分比':<20}"
            txt.append(wait_class_header)
            txt.append("-" * 50)
            
            for wait_class, pct in wait_class_distribution.items():
                wait_class_row = f"{wait_class:<30} {str(pct)+'%':<20}"
                txt.append(wait_class_row)
        else:
            txt.append("没有等待类别分布数据")
        
        txt.append("")
        txt.append("6.3 等待事件特征")
        txt.append("")
        txt.append(f"CPU绑定: {'是' if wait_analysis.get('cpu_bound', False) else '否'}")
        txt.append(f"IO绑定: {'是' if wait_analysis.get('io_bound', False) else '否'}")
        txt.append(f"并发问题: {'是' if wait_analysis.get('concurrency_issues', False) else '否'}")
        txt.append(f"网络问题: {'是' if wait_analysis.get('network_issues', False) else '否'}")
        txt.append("")
        
        # 添加潜在问题
        txt.append("7. 潜在问题")
        txt.append("-" * 80)
        potential_issues = self.analysis_results.get('potential_issues', [])
        
        if potential_issues:
            for i, issue in enumerate(potential_issues, 1):
                txt.append(f"7.{i} {issue.get('issue', 'Unknown')} ({issue.get('severity', 'Unknown')})")
                txt.append("")
                txt.append(f"类别: {issue.get('category', '')}")
                txt.append(f"描述: {issue.get('description', '')}")
                txt.append(f"建议: {issue.get('recommendation', '')}")
                txt.append("")
        else:
            txt.append("没有发现潜在问题")
            txt.append("")
        
        # 添加结论
        txt.append("8. 结论")
        txt.append("-" * 80)
        performance_score = summary.get('performance_score', 0)
        performance_level = summary.get('performance_level', 'Unknown')
        
        txt.append(f"根据AWR报告分析，数据库性能评分为{performance_score}/100，性能等级为{performance_level}。")
        txt.append(f"{summary.get('overall_assessment', '')}")
        
        if performance_score < 60:
            txt.append("建议优先解决上述优先建议中的高严重性问题，以提高数据库性能。")
        elif performance_score < 80:
            txt.append("建议按照优先级顺序解决上述问题，以进一步提高数据库性能。")
        else:
            txt.append("数据库性能良好，建议定期监控并进行预防性维护。")
        
        # 返回文本内容
        return "\n".join(txt)
    
    def _generate_doc_content(self, doc):
        """生成Word文档内容"""
        # 设置样式
        styles = doc.styles
        
        # 添加标题
        summary = self.recommendations.get('summary', {})
        db_name = summary.get('database_name', 'Unknown')
        instance_name = summary.get('instance_name', 'Unknown')
        
        title = doc.add_heading(f"Oracle AWR报告分析 - {db_name}/{instance_name}", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph(f"生成时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph("由青学会MOP技术社区AWR分析工具生成")
        
        # 添加概要信息
        doc.add_heading("1. 分析概要", 1)
        
        summary_table = doc.add_table(rows=7, cols=2)
        summary_table.style = 'Table Grid'
        
        # 填充概要表格
        summary_rows = [
            ["数据库名称", summary.get('database_name', 'Unknown')],
            ["实例名称", summary.get('instance_name', 'Unknown')],
            ["数据库版本", summary.get('version', 'Unknown')],
            ["主机名", summary.get('host', 'Unknown')],
            ["快照时间段", summary.get('snapshot_period', 'Unknown')],
            ["性能评分", f"{summary.get('performance_score', 0)}/100 ({summary.get('performance_level', 'Unknown')})"],
            ["整体评估", summary.get('overall_assessment', '')]
        ]
        
        for i, row_data in enumerate(summary_rows):
            row = summary_table.rows[i]
            row.cells[0].text = row_data[0]
            row.cells[1].text = row_data[1]
            row.cells[0].width = Inches(2)
            row.cells[1].width = Inches(4)
        
        # 添加优先建议
        doc.add_heading("2. 优先建议", 1)
        priority_recommendations = self.recommendations.get('priority_recommendations', [])
        
        if priority_recommendations:
            for i, rec in enumerate(priority_recommendations, 1):
                doc.add_heading(f"2.{i} {rec.get('title', 'Unknown')} ({rec.get('severity', 'Unknown')})", 2)
                
                doc.add_paragraph(f"发现: {rec.get('finding', '')}")
                doc.add_paragraph(f"影响: {rec.get('impact', '')}")
                
                p = doc.add_paragraph("建议:")
                for r in rec.get('recommendations', []):
                    p = doc.add_paragraph(r)
                    p.style = 'List Bullet'
        else:
            doc.add_paragraph("没有优先建议")
        
        # 添加CPU分析
        doc.add_heading("3. CPU分析", 1)
        cpu_analysis = self.analysis_results.get('cpu_analysis', {})
        
        cpu_table = doc.add_table(rows=9, cols=2)
        cpu_table.style = 'Table Grid'
        
        # 填充CPU表格
        cpu_rows = [
            ["CPU核心数", str(cpu_analysis.get('cpu_cores', 0))],
            ["主机用户CPU", f"{cpu_analysis.get('host_cpu_usage', {}).get('user', 0)}%"],
            ["主机系统CPU", f"{cpu_analysis.get('host_cpu_usage', {}).get('system', 0)}%"],
            ["主机空闲CPU", f"{cpu_analysis.get('host_cpu_usage', {}).get('idle', 0)}%"],
            ["主机总CPU使用率", f"{cpu_analysis.get('host_cpu_usage', {}).get('total', 0)}%"],
            ["数据库CPU占用", f"{cpu_analysis.get('db_cpu_usage', {}).get('total_cpu_pct', 0)}%"],
            ["每秒CPU时间", f"{cpu_analysis.get('db_cpu_usage', {}).get('per_second', 0)}秒"],
            ["CPU利用率", f"{cpu_analysis.get('cpu_utilization', 0)}%"],
            ["CPU瓶颈", "是" if cpu_analysis.get('cpu_bottleneck', False) else "否"]
        ]
        
        for i, row_data in enumerate(cpu_rows):
            row = cpu_table.rows[i]
            row.cells[0].text = row_data[0]
            row.cells[1].text = row_data[1]
            row.cells[0].width = Inches(2)
            row.cells[1].width = Inches(4)
        
        # 添加内存分析
        doc.add_heading("4. 内存分析", 1)
        memory_analysis = self.analysis_results.get('memory_analysis', {})
        
        memory_table = doc.add_table(rows=6, cols=2)
        memory_table.style = 'Table Grid'
        
        # 填充内存表格
        memory_rows = [
            ["缓冲区命中率", f"{memory_analysis.get('buffer_hit_ratio', 0)}%"],
            ["库缓存命中率", f"{memory_analysis.get('library_hit_ratio', 0)}%"],
            ["内存排序百分比", f"{memory_analysis.get('in_memory_sort_pct', 0)}%"],
            ["低缓冲区命中", "是" if memory_analysis.get('memory_issues', {}).get('low_buffer_hit', False) else "否"],
            ["低库缓存命中", "是" if memory_analysis.get('memory_issues', {}).get('low_library_hit', False) else "否"],
            ["低内存排序", "是" if memory_analysis.get('memory_issues', {}).get('low_memory_sort', False) else "否"]
        ]
        
        for i, row_data in enumerate(memory_rows):
            row = memory_table.rows[i]
            row.cells[0].text = row_data[0]
            row.cells[1].text = row_data[1]
            row.cells[0].width = Inches(2)
            row.cells[1].width = Inches(4)
        
        # 添加IO分析
        doc.add_heading("5. IO分析", 1)
        io_analysis = self.analysis_results.get('io_analysis', {})
        
        io_table = doc.add_table(rows=7, cols=2)
        io_table.style = 'Table Grid'
        
        # 填充IO表格
        io_rows = [
            ["每秒物理读", str(io_analysis.get('physical_reads_per_sec', 0))],
            ["每秒物理写", str(io_analysis.get('physical_writes_per_sec', 0))],
            ["每秒读IO请求", str(io_analysis.get('read_io_requests_per_sec', 0))],
            ["每秒写IO请求", str(io_analysis.get('write_io_requests_per_sec', 0))],
            ["每秒总IO请求", str(io_analysis.get('total_io_requests_per_sec', 0))],
            ["读写比例", f"{io_analysis.get('read_write_ratio', 0)}:1"],
            ["IO瓶颈", "是" if io_analysis.get('io_bottleneck', False) else "否"]
        ]
        
        for i, row_data in enumerate(io_rows):
            row = io_table.rows[i]
            row.cells[0].text = row_data[0]
            row.cells[1].text = row_data[1]
            row.cells[0].width = Inches(2)
            row.cells[1].width = Inches(4)
        
        # 添加等待事件分析
        doc.add_heading("6. 等待事件分析", 1)
        wait_analysis = self.analysis_results.get('wait_analysis', {})
        
        doc.add_heading("6.1 前5个耗时事件", 2)
        top_events = wait_analysis.get('top_events', [])
        
        if top_events:
            event_table = doc.add_table(rows=len(top_events)+1, cols=4)
            event_table.style = 'Table Grid'
            
            # 添加表头
            header_row = event_table.rows[0]
            header_row.cells[0].text = "事件名称"
            header_row.cells[1].text = "等待时间(秒)"
            header_row.cells[2].text = "数据库时间百分比"
            header_row.cells[3].text = "等待类别"
            
            # 填充事件表格
            for i, event in enumerate(top_events, 1):
                row = event_table.rows[i]
                row.cells[0].text = event.get('event_name', '')
                row.cells[1].text = str(event.get('wait_time', 0))
                row.cells[2].text = f"{event.get('db_time_pct', 0)}%"
                row.cells[3].text = event.get('wait_class', '')
        else:
            doc.add_paragraph("没有等待事件数据")
        
        doc.add_heading("6.2 等待类别分布", 2)
        wait_class_distribution = wait_analysis.get('wait_class_distribution', {})
        
        if wait_class_distribution:
            wait_class_table = doc.add_table(rows=len(wait_class_distribution)+1, cols=2)
            wait_class_table.style = 'Table Grid'
            
            # 添加表头
            header_row = wait_class_table.rows[0]
            header_row.cells[0].text = "等待类别"
            header_row.cells[1].text = "数据库时间百分比"
            
            # 填充等待类别表格
            for i, (wait_class, pct) in enumerate(wait_class_distribution.items(), 1):
                row = wait_class_table.rows[i]
                row.cells[0].text = wait_class
                row.cells[1].text = f"{pct}%"
        else:
            doc.add_paragraph("没有等待类别分布数据")
        
        doc.add_heading("6.3 等待事件特征", 2)
        
        wait_feature_table = doc.add_table(rows=4, cols=2)
        wait_feature_table.style = 'Table Grid'
        
        # 填充等待特征表格
        wait_feature_rows = [
            ["CPU绑定", "是" if wait_analysis.get('cpu_bound', False) else "否"],
            ["IO绑定", "是" if wait_analysis.get('io_bound', False) else "否"],
            ["并发问题", "是" if wait_analysis.get('concurrency_issues', False) else "否"],
            ["网络问题", "是" if wait_analysis.get('network_issues', False) else "否"]
        ]
        
        for i, row_data in enumerate(wait_feature_rows):
            row = wait_feature_table.rows[i]
            row.cells[0].text = row_data[0]
            row.cells[1].text = row_data[1]
            row.cells[0].width = Inches(2)
            row.cells[1].width = Inches(4)
        
        # 添加结论
        doc.add_heading("7. 结论", 1)
        performance_score = summary.get('performance_score', 0)
        performance_level = summary.get('performance_level', 'Unknown')
        
        conclusion = doc.add_paragraph()
        conclusion.add_run(f"根据AWR报告分析，数据库性能评分为{performance_score}/100，性能等级为{performance_level}。")
        conclusion.add_run(f" {summary.get('overall_assessment', '')}")
        
        if performance_score < 60:
            conclusion.add_run(" 建议优先解决上述优先建议中的高严重性问题，以提高数据库性能。")
        elif performance_score < 80:
            conclusion.add_run(" 建议按照优先级顺序解决上述问题，以进一步提高数据库性能。")
        else:
            conclusion.add_run(" 数据库性能良好，建议定期监控并进行预防性维护。")
